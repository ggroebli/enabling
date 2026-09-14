# -*- coding: utf-8 -*-
"""
DUB Verkaeufer-Academy - Bewertungs-Template (Modul 4)
Erzeugt Downloads/Verkaeufer/bewertungs-template.xlsx

Aufbau je Blatt:
  1 Zahlen aus dem Jahresabschluss  -> EBITDA laut Jahresabschluss
  2 Bereinigungen                   -> Bereinigtes EBITDA
  3 Gewichtung der Jahre            -> Gewichtetes bereinigtes EBITDA
  4 EBITDA-Multiples der Branche    -> Min / Durchschnitt / Max
  5 Bewertungsbandbreite            -> Unternehmenswert
  6 Wo in der Spanne liegt dein Unternehmen?
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- CI-Tokens
NAVY       = "0F2744"
NAVY2      = "14294A"
ORANGE     = "EE7A1E"
ORANGE_LT  = "FDEBD8"
INK        = "1A1A1A"
GREY       = "666666"
GREY_LT    = "8A94A6"
LINE       = "E5E7EB"
PAPER      = "EEF1F5"
INPUT_BG   = "FFF8EF"   # Eingabefelder: heller Orangeton
INPUT_FG   = "B35C0F"   # Eingabewerte: dunkles Orange
WHITE      = "FFFFFF"
NAVY_TEXT  = "C7D0DE"

FONT = "Calibri"

EUR  = '#,##0 "€";[Red]-#,##0 "€"'
EURC = '#,##0 "€";[Red]-#,##0 "€";""'      # berechnet: Null bleibt leer
PCT = '0 %'
MUL = '0.0"×"'
YR  = '0'


def f(sz=10.5, b=False, color=INK, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
LEFTW  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LABELW = Alignment(horizontal="left",   vertical="center", wrap_text=True)

COLS = ["C", "D", "E", "F"]          # die vier Jahresspalten
MULCOLS = ["D", "E", "F"]            # Min / Durchschnitt / Max


# ---------------------------------------------------------------- Beispiel
BEISPIEL = {
    "jahre":            [2024, 2025, 2026, 2027],
    "umsatz":           [720000, 780000, 830000, 880000],
    "gesamtleistung":   [730000, 790000, 840000, 890000],
    "egt":              [104000, 112000, 120000, 131000],
    "finanzergebnis":   [-12000, -11000, -10000, -9000],
    "afa":              [32000, 32000, 32000, 32000],
    "ber_gehalt":       [30000, 30000, 32000, 32000],
    "ber_privat":       [6000, 6000, 6000, 6000],
    "ber_aufwand":      [0, 0, 0, 0],
    "ber_ertrag":       [0, 0, 0, 0],
    "ber_miete":        [0, 0, 0, 0],
    "ber_frei":         [0, 0, 0, 0],
    "multiples":        [3.5, 4.0, 5.0],
}
LEER = {k: ([None] * 4 if k != "multiples" else [None] * 3) for k in BEISPIEL}
LEER["jahre"] = [2024, 2025, 2026, 2027]


FAKTOREN = [
    ("Exit-Readiness und Inhaberabhängigkeit",
     "Oberer Rand: Das Tagesgeschäft läuft ohne dich, Prozesse sind dokumentiert, Kundenbeziehungen hängen am Team.   "
     "Unterer Rand: Entscheidungen, Fachwissen und Kundenkontakte liegen allein bei dir."),
    ("Wachstum",
     "Oberer Rand: Umsatz und Ergebnis sind über mehrere Jahre gestiegen, die Planung ist mit Aufträgen unterlegt.   "
     "Unterer Rand: Rückläufige oder stark schwankende Zahlen ohne erkennbaren Grund."),
    ("Marge im Branchenvergleich",
     "Oberer Rand: Deine bereinigte EBITDA-Marge liegt über dem Branchenschnitt.   "
     "Unterer Rand: Sie liegt darunter, ohne dass es dafür eine erklärbare Sondersituation gibt."),
    ("Personalstruktur",
     "Oberer Rand: Eingespieltes Team, Schlüsselpositionen doppelt besetzt, geringe Fluktuation.   "
     "Unterer Rand: Offene Schlüsselstellen, hohe Fluktuation oder ein Team kurz vor dem Ruhestand."),
    ("Kundenstruktur",
     "Oberer Rand: Viele Kunden, kein einzelner dominiert, wiederkehrende Erlöse aus Wartung oder Rahmenverträgen.   "
     "Unterer Rand: Ein einzelner Kunde macht 30 % oder mehr des Umsatzes aus."),
    ("Verträge und Ausstattung",
     "Oberer Rand: Langfristige Miet-, Liefer- und Kundenverträge, gepflegte Maschinen und Software.   "
     "Unterer Rand: Kurzfristig kündbare Verträge, Investitionsstau oder ungeklärte Rechtsfragen."),
]


class Sheet:
    """Schreibt das Modell zeilenweise und merkt sich die Zeilennummern."""

    def __init__(self, ws, data, beispiel=False):
        self.ws = ws
        self.d = data
        self.beispiel = beispiel
        self.r = 1
        self.ref = {}

    # -- Bausteine ------------------------------------------------
    def row_h(self, h):
        self.ws.row_dimensions[self.r].height = h

    def band(self, h):
        """volle Navy-Zeile"""
        for col in "ABCDEFG":
            self.ws[f"{col}{self.r}"].fill = fill(NAVY)
        self.row_h(h)
        self.r += 1

    def section(self, nr, titel):
        ws = self.ws
        self.row_h(7)
        self.r += 1
        ws.merge_cells(f"B{self.r}:F{self.r}")
        c = ws[f"B{self.r}"]
        c.value = f"  {nr}   {titel.upper()}"
        c.font = f(11, True, WHITE)
        c.alignment = LEFT
        for col in "BCDEF":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(24)
        self.r += 1

    def note(self, text, color=GREY, size=9.5, height=None, italic=False):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:F{self.r}")
        c = ws[f"B{self.r}"]
        c.value = text
        c.font = f(size, False, color, italic)
        c.alignment = LEFTW
        self.row_h(14)   # AutoFit vergroessert bei Bedarf
        self.r += 1

    def input_row(self, label, key, fmt=EUR, indent=False, height=17):
        """Eingabezeile mit vier Jahresspalten."""
        ws = self.ws
        c = ws[f"B{self.r}"]
        c.value = ("      " if indent else "") + label
        c.font = f(10.5)
        c.alignment = LABELW
        c.border = Border(bottom=side(LINE))
        vals = self.d.get(key, [None] * 4)
        for i, col in enumerate(COLS):
            cell = ws[f"{col}{self.r}"]
            cell.value = vals[i]
            cell.number_format = fmt
            cell.font = f(10.5, False, INPUT_FG)
            cell.fill = fill(INPUT_BG)
            cell.alignment = RIGHT
            cell.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        self.row_h(height)
        row = self.r
        self.r += 1
        return row

    def calc_row(self, label, formula_tpl, fmt=EURC, strong=False, accent=False,
                 size=10.5, height=19):
        """Berechnete Zeile; formula_tpl bekommt die Spalte als {c}."""
        ws = self.ws
        bg = ORANGE_LT if accent else PAPER
        c = ws[f"B{self.r}"]
        c.value = label
        if isinstance(label, str) and label.startswith("="):
            c.data_type = "s"          # sonst liest Excel das Label als Formel
        c.font = f(size, True, NAVY if accent else INK)
        c.alignment = LABELW
        c.fill = fill(bg)
        c.border = Border(top=side(NAVY if accent else GREY_LT), bottom=side(LINE))
        for col in COLS:
            cell = ws[f"{col}{self.r}"]
            cell.value = formula_tpl.format(c=col)
            cell.number_format = fmt
            cell.font = f(size, True, NAVY if accent else INK)
            cell.fill = fill(bg)
            cell.alignment = RIGHT
            cell.border = Border(top=side(NAVY if accent else GREY_LT), bottom=side(LINE))
        self.row_h(height)
        row = self.r
        self.r += 1
        return row

    def spacer(self, h=8):
        self.row_h(h)
        self.r += 1

    # -- Abschnitte -----------------------------------------------
    def kopf(self):
        ws = self.ws
        self.band(6)
        # Zeile Marke
        ws[f"B{self.r}"] = "DUB  ·  VERKÄUFER-ACADEMY"
        ws[f"B{self.r}"].font = f(10, True, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        ws.merge_cells(f"D{self.r}:F{self.r}")
        ws[f"D{self.r}"] = "MODUL 4 · WAS IST MEIN UNTERNEHMEN WERT?"
        ws[f"D{self.r}"].font = f(10, True, ORANGE)
        ws[f"D{self.r}"].alignment = RIGHT
        self.band(18)

        ws.merge_cells(f"B{self.r}:F{self.r}")
        titel = "Bewertungs-Template" + ("  –  ausgefülltes Beispiel" if self.beispiel else "")
        ws[f"B{self.r}"] = titel
        ws[f"B{self.r}"].font = f(20, True, WHITE)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(30)

        ws.merge_cells(f"B{self.r}:F{self.r}")
        ws[f"B{self.r}"] = "Bereinigtes EBITDA × Branchen-Multiple = deine Bewertungsbandbreite"
        ws[f"B{self.r}"].font = f(11, False, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(19)
        self.band(8)
        self.spacer(10)

    def anleitung(self):
        if self.beispiel:
            self.ws.merge_cells(f"B{self.r}:F{self.r}")
            c = self.ws[f"B{self.r}"]
            c.value = ("Dieses Blatt zeigt das Modell mit Beispielzahlen eines Handwerksbetriebs. "
                       "Es dient nur der Veranschaulichung – arbeite im Blatt „Bewertung“ mit deinen eigenen Zahlen.")
            c.font = f(10, False, NAVY)
            c.fill = fill(ORANGE_LT)
            c.alignment = LEFTW
            for col in "BCDEF":
                self.ws[f"{col}{self.r}"].fill = fill(ORANGE_LT)
            self.row_h(30)
            self.r += 1
            self.spacer(10)
            return

        self.section("▶", "So gehst du vor")
        for t in [
            "1.  Trage in Abschnitt 1 die Zahlen aus deinem Jahresabschluss ein: für die beiden letzten abgeschlossenen Jahre, "
            "deine Hochrechnung für das laufende Jahr und deine Planung für das Folgejahr.",
            "2.  Ergänze in Abschnitt 2 die Bereinigungen. Sie machen sichtbar, was dein Unternehmen ohne inhaberbedingte "
            "Sondereffekte tatsächlich verdient – die wichtigste Kennzahl für Käufer.",
            "3.  Trage in Abschnitt 4 die EBITDA-Multiples für deine Branche ein. Die aktuellen KMU-Multiples findest du auf www.dub.de.",
            "4.  Abschnitt 5 zeigt deine Bewertungsbandbreite, Abschnitt 6 hilft dir einzuordnen, wo in dieser Spanne dein Unternehmen steht.",
        ]:
            self.note(t, INK, 10, 15)
        self.note("Die orange hinterlegten Felder sind Eingabefelder. Alle übrigen Werte berechnen sich automatisch.",
                  GREY_LT, 9.5, 16, italic=True)
        self.spacer(10)

    def abschnitt_1(self):
        ws = self.ws
        self.section("1", "Zahlen aus deinem Jahresabschluss")
        # Spaltenkopf
        heads = ["Zwei Jahre zurück\n(abgeschlossen)", "Letztes Jahr\n(abgeschlossen)",
                 "Laufendes Jahr\n(Forecast)", "Folgejahr\n(Plan)"]
        ws[f"B{self.r}"] = "Position · alle Beträge in Euro"
        ws[f"B{self.r}"].font = f(9.5, True, GREY)
        ws[f"B{self.r}"].alignment = Alignment(horizontal="left", vertical="bottom")
        for col, h in zip(COLS, heads):
            c = ws[f"{col}{self.r}"]
            c.value = h
            c.font = f(9, True, NAVY)
            c.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
        self.row_h(30)
        self.r += 1

        self.ref["jahr"] = self.input_row("Geschäftsjahr", "jahre", YR)
        self.ref["umsatz"] = self.input_row("Umsatzerlöse", "umsatz")
        self.input_row("Gesamtleistung (Umsatz ± Bestandsveränderung + aktivierte Eigenleistungen)",
                       "gesamtleistung")
        self.ref["gl"] = self.r - 1
        self.input_row("Ergebnis der gewöhnlichen Geschäftstätigkeit (EGT)", "egt")
        self.ref["egt"] = self.r - 1
        self.input_row("Finanzergebnis (Zinserträge − Zinsaufwendungen, meist negativ)", "finanzergebnis")
        self.ref["fe"] = self.r - 1
        self.input_row("Abschreibungen auf Sachanlagen und immaterielle Vermögensgegenstände", "afa")
        self.ref["afa"] = self.r - 1

        self.ref["ebitda"] = self.calc_row(
            "= EBITDA laut Jahresabschluss",
            "={{c}}{egt}-{{c}}{fe}+{{c}}{afa}".format(**self.ref))
        self.note("EBITDA = EGT abzüglich Finanzergebnis zuzüglich Abschreibungen. Weist deine GuV kein EGT aus, "
                  "nimm stattdessen: Jahresüberschuss + Steuern vom Einkommen und Ertrag.", GREY_LT, 9.5, 15, italic=True)
        self.spacer(10)

    def abschnitt_2(self):
        self.section("2", "Bereinigungen")
        self.note("Bei inhabergeführten Unternehmen muss das EBITDA bereinigt werden, damit Käufer die tatsächliche "
                  "Ertragskraft erkennen. Alles, was das Ergebnis nach der Bereinigung erhöht, trägst du positiv ein, "
                  "alles Mindernde mit Minuszeichen.", INK, 10, 28)
        start = self.r
        self.input_row("Inhabergehalt: Differenz zu einem marktüblichen Geschäftsführergehalt", "ber_gehalt", indent=True)
        self.input_row("Private Ausgaben über das Unternehmen (z. B. privat genutzter Firmenwagen)", "ber_privat", indent=True)
        self.input_row("Einmalige Aufwendungen (z. B. Rechtsstreit, Umzug, Abfindung)", "ber_aufwand", indent=True)
        self.input_row("Einmalige Erträge (z. B. Versicherungsleistung, Anlagenverkauf) – negativ eintragen", "ber_ertrag", indent=True)
        self.input_row("Miete an dich oder verbundene Unternehmen: Differenz zur marktüblichen Miete", "ber_miete", indent=True)
        self.input_row("Weitere Bereinigung (frei)", "ber_frei", indent=True)
        ende = self.r - 1

        self.ref["summe_ber"] = self.calc_row(
            "= Summe Bereinigungen", "=SUM({{c}}{a}:{{c}}{b})".format(a=start, b=ende))
        self.ref["ber_ebitda"] = self.calc_row(
            "= BEREINIGTES EBITDA", "={{c}}{e}+{{c}}{s}".format(e=self.ref["ebitda"], s=self.ref["summe_ber"]),
            accent=True, size=12, height=24)
        self.calc_row("Bereinigte EBITDA-Marge (auf Gesamtleistung)",
                      '=IFERROR({{c}}{b}/{{c}}{g},"")'.format(b=self.ref["ber_ebitda"], g=self.ref["gl"]),
                      fmt=PCT, size=10, height=17)
        self.note("Käufer hinterfragen jede Bereinigung. Je nachvollziehbarer du sie belegen kannst, desto stärker "
                  "deine Verhandlungsposition.", GREY_LT, 9.5, 15, italic=True)
        self.spacer(10)

    def abschnitt_3(self):
        ws = self.ws
        self.section("3", "Gewichtung der Jahre")
        self.note("Käufer schauen auf die belegte Vergangenheit und auf die Perspektive. Das letzte abgeschlossene Jahr und "
                  "das laufende Jahr wiegen deshalb am schwersten. Du kannst die Gewichtung anpassen, in Summe muss sie 100 % ergeben.",
                  INK, 10, 28)
        gew = [0.15, 0.35, 0.35, 0.15]
        c = ws[f"B{self.r}"]
        c.value = "Gewichtung"
        c.font = f(10.5)
        c.alignment = LABELW
        c.border = Border(bottom=side(LINE))
        for col, g in zip(COLS, gew):
            cell = ws[f"{col}{self.r}"]
            cell.value = g
            cell.number_format = PCT
            cell.font = f(10.5, True, INPUT_FG)
            cell.fill = fill(INPUT_BG)
            cell.alignment = RIGHT
            cell.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        self.ref["gew"] = self.r
        self.row_h(19)
        self.r += 1

        ws.merge_cells(f"B{self.r}:F{self.r}")
        chk = ws[f"B{self.r}"]
        chk.value = ('=IF(SUM(C{g}:F{g})=1,"Gewichtung ergibt 100 % – passt.",'
                     '"Achtung: Die Gewichtung ergibt "&TEXT(SUM(C{g}:F{g}),"0 %")&". Bitte auf 100 % anpassen.")'
                     ).format(g=self.ref["gew"])
        chk.font = f(9.5, False, GREY_LT, True)
        chk.alignment = LEFTW
        self.row_h(15)
        self.r += 1
        self.spacer(4)

        # Ergebniszeile: Label B:D, Wert E:F
        ws.merge_cells(f"B{self.r}:D{self.r}")
        lab = ws[f"B{self.r}"]
        lab.value = "= GEWICHTETES BEREINIGTES EBITDA  ·  deine Bewertungsbasis"
        lab.data_type = "s"
        lab.font = f(12, True, NAVY)
        lab.alignment = LEFT
        ws.merge_cells(f"E{self.r}:F{self.r}")
        val = ws[f"E{self.r}"]
        val.value = "=SUMPRODUCT(C{b}:F{b},C{g}:F{g})".format(b=self.ref["ber_ebitda"], g=self.ref["gew"])
        val.number_format = EURC
        val.font = f(13, True, NAVY)
        val.alignment = RIGHT
        for col in "BCDEF":
            ws[f"{col}{self.r}"].fill = fill(ORANGE_LT)
            ws[f"{col}{self.r}"].border = Border(top=side(NAVY), bottom=side(NAVY))
        self.ref["basis"] = self.r
        self.row_h(26)
        self.r += 1
        self.spacer(10)

    def abschnitt_4(self):
        ws = self.ws
        self.section("4", "EBITDA-Multiples deiner Branche")
        self.note("Die aktuellen KMU-Multiples nach Branche findest du auf www.dub.de. Trage die drei Werte für deine "
                  "Branche ein – sie bestimmen, wie breit deine Bewertungsspanne ausfällt.", INK, 10, 28)
        for col, h in zip(MULCOLS, ["Minimum", "Durchschnitt", "Maximum"]):
            c = ws[f"{col}{self.r}"]
            c.value = h
            c.font = f(9.5, True, NAVY)
            c.alignment = CENTER
        self.row_h(18)
        self.r += 1

        c = ws[f"B{self.r}"]
        c.value = "EBITDA-Multiple deiner Branche"
        c.font = f(10.5)
        c.alignment = LABELW
        c.border = Border(bottom=side(LINE))
        for col, v in zip(MULCOLS, self.d["multiples"]):
            cell = ws[f"{col}{self.r}"]
            cell.value = v
            cell.number_format = MUL
            cell.font = f(11, True, INPUT_FG)
            cell.fill = fill(INPUT_BG)
            cell.alignment = RIGHT
            cell.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        self.ref["mult"] = self.r
        self.row_h(21)
        self.r += 1
        self.spacer(10)

    def abschnitt_5(self):
        ws = self.ws
        self.section("5", "Deine Bewertungsbandbreite")
        for col, h in zip(MULCOLS, ["Minimum", "Durchschnitt", "Maximum"]):
            c = ws[f"{col}{self.r}"]
            c.value = h
            c.font = f(9.5, True, NAVY)
            c.alignment = CENTER
        self.row_h(18)
        self.r += 1

        c = ws[f"B{self.r}"]
        ws.merge_cells(f"B{self.r}:C{self.r}")
        c.value = "UNTERNEHMENSWERT"
        c.font = f(13, True, WHITE)
        c.alignment = LEFT
        for col in MULCOLS:
            cell = ws[f"{col}{self.r}"]
            cell.value = "=$E${b}*{m}{r}".format(b=self.ref["basis"], m=col, r=self.ref["mult"])
            cell.number_format = EURC
            cell.font = f(13, True, WHITE)
            cell.alignment = RIGHT
        for col in "BCDEF":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(32)
        self.r += 1

        self.note("Das ist der Unternehmenswert, nicht der Kaufpreis. Er ist „cash and debt free“ gerechnet: "
                  "Finanzverbindlichkeiten, überschüssige Barmittel, Working Capital und Investitionsstau kommen erst "
                  "über die Preisbrücke dazu oder gehen ab (siehe Modul 4, Abschnitt 4).", GREY, 9.5, 28)
        self.spacer(10)

    def abschnitt_6(self):
        ws = self.ws
        self.ref["break6"] = self.r
        self.section("6", "Wo in dieser Spanne liegt dein Unternehmen?")
        ws.merge_cells(f"B{self.r}:F{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("Ob dein Unternehmen eher am unteren oder am oberen Rand dieser Spanne liegt, entscheidet sich daran, "
                   "wie gut es ohne dich funktioniert und wie verlässlich seine Erträge sind: an Exit-Readiness und "
                   "Inhaberabhängigkeit, am Wachstumstrend, an der Marge im Branchenvergleich, an der Personalstruktur, "
                   "an der Kundenstruktur und an der Qualität deiner Verträge.")
        c.font = f(11, True, NAVY)
        c.alignment = LEFTW
        for col in "BCDEF":
            ws[f"{col}{self.r}"].fill = fill(ORANGE_LT)
        self.row_h(46)
        self.r += 1
        self.spacer(6)

        for name, text in FAKTOREN:
            b = ws[f"B{self.r}"]
            b.value = name
            b.font = f(10, True, NAVY)
            b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            b.border = Border(top=side(LINE))
            ws.merge_cells(f"C{self.r}:F{self.r}")
            t = ws[f"C{self.r}"]
            t.value = text
            t.font = f(9.5, False, INK)
            t.alignment = LEFTW
            for col in "CDEF":
                ws[f"{col}{self.r}"].border = Border(top=side(LINE))
            self.row_h(30)
            self.r += 1
        self.spacer(12)

    def fuss(self):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:F{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("Dieses Template dient der Orientierung und ersetzt keine Rechts-, Steuer- oder Finanzberatung.    "
                   "DUB · info@dub.de · www.dub.de    © 2026 Deutsche Unternehmerbörse DUB.de GmbH")
        c.font = f(9, False, GREY_LT)
        c.alignment = LEFTW
        for col in "BCDEF":
            ws[f"{col}{self.r}"].border = Border(top=side(LINE))
        self.row_h(24)
        self.r += 1

    def layout(self):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = NAVY
        widths = {"A": 2.2, "B": 74, "C": 16, "D": 16, "E": 16, "F": 16, "G": 2.2}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4

    def build(self):
        self.layout()
        self.kopf()
        self.anleitung()
        self.abschnitt_1()
        self.abschnitt_2()
        self.abschnitt_3()
        self.abschnitt_4()
        self.abschnitt_5()
        self.abschnitt_6()
        self.fuss()
        # Abschnitt 6 soll nicht am Seitenfuss angerissen werden
        self.ws.row_breaks.append(Break(id=self.ref["break6"]))


# ------------------------------------------------------------------ AutoFit
# Excel misst die noetigen Zeilenhoehen selbst. Verbundene Zellen ignoriert
# Excels AutoFit, deshalb wird ihr Text in eine Hilfsspalte gleicher Breite
# gespiegelt und die Zeile darueber gemessen.
AUTOFIT_PS = r"""
param([string]$Path)
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
$wb = $xl.Workbooks.Open($Path)
$HELP = 26
foreach ($ws in $wb.Worksheets) {
    $used = $ws.UsedRange
    $last = $used.Row + $used.Rows.Count - 1
    $minH = @{}
    for ($r = 1; $r -le $last; $r++) { $minH[$r] = $ws.Rows.Item($r).RowHeight }
    for ($r = 1; $r -le $last; $r++) {
        $text = $null; $fromCol = 0; $span = 0
        for ($c = 2; $c -le 6; $c++) {
            $cell = $ws.Cells.Item($r, $c)
            if ($cell.MergeCells -and $cell.WrapText) {
                $ma = $cell.MergeArea
                if ($ma.Rows.Count -eq 1 -and $ma.Row -eq $r) {
                    $tl = $ws.Cells.Item($r, $ma.Column)
                    if ($tl.Value2 -is [string] -and $tl.Value2.Length -gt 0) {
                        $text = $tl.Value2; $fromCol = $ma.Column; $span = $ma.Columns.Count
                    }
                }
                break
            }
        }
        if ($null -ne $text) {
            $w = 0
            for ($k = 0; $k -lt $span; $k++) { $w += $ws.Columns.Item($fromCol + $k).ColumnWidth }
            $h = $ws.Cells.Item($r, $fromCol)
            $ws.Columns.Item($HELP).ColumnWidth = $w
            $probe = $ws.Cells.Item($r, $HELP)
            $probe.Font.Name = $h.Font.Name; $probe.Font.Size = $h.Font.Size
            $probe.Font.Bold = $h.Font.Bold; $probe.WrapText = $true
            $probe.Value2 = $text
        }
        $fit = ($null -ne $text)
        if (-not $fit) {
            for ($c = 2; $c -le 6; $c++) {
                $cell = $ws.Cells.Item($r, $c)
                if ((-not $cell.MergeCells) -and $cell.WrapText -and $null -ne $cell.Value2) { $fit = $true; break }
            }
        }
        if ($fit) {
            $ws.Rows.Item($r).AutoFit() | Out-Null
            $ws.Rows.Item($r).RowHeight = [Math]::Max($ws.Rows.Item($r).RowHeight + 3, $minH[$r])
            if ($null -ne $text) { $ws.Cells.Item($r, $HELP).Clear() | Out-Null }
        }
    }
    $ws.Columns.Item($HELP).ColumnWidth = 8.43
}
$wb.Worksheets.Item(1).Activate()
$wb.Save(); $wb.Close($true); $xl.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null
"""


def autofit(path):
    """Zeilenhoehen von Excel messen lassen (nur Windows mit installiertem Excel)."""
    import subprocess, tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as fh:
        fh.write(AUTOFIT_PS)
        script = fh.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", script, "-Path", os.path.abspath(path)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            print("AutoFit uebersprungen (kein Excel?):", (r.stderr or "").strip()[:200])
        else:
            print("Zeilenhoehen von Excel gemessen.")
    finally:
        os.unlink(script)


def main():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Bewertung"
    Sheet(ws1, LEER, beispiel=False).build()

    ws2 = wb.create_sheet("Beispiel")
    Sheet(ws2, BEISPIEL, beispiel=True).build()
    ws2.sheet_properties.tabColor = ORANGE

    wb.properties.title = "DUB Bewertungs-Template"
    wb.properties.creator = "Deutsche Unternehmerbörse DUB.de GmbH"
    wb.properties.description = "Verkäufer-Academy, Modul 4: Was ist mein Unternehmen wert?"

    out = os.path.join("Downloads", "Verkaeufer", "bewertungs-template.xlsx")
    wb.save(out)
    print("geschrieben:", out)
    autofit(out)


if __name__ == "__main__":
    main()
