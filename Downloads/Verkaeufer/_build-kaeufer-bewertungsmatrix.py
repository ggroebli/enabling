# -*- coding: utf-8 -*-
"""
DUB Verkaeufer-Academy - Kaeufer-Bewertungsmatrix (Modul 6)
Erzeugt Downloads/Verkaeufer/kaeufer-bewertungsmatrix.xlsx

Inhalt kongruent zu Modul 6 "Kaeufer- und Angebots-Screening" und zur
Kaeufer-Screening-Checkliste. Layout wie bewertungs-template.xlsx.

Aufbau je Blatt:
  1 Die Kandidaten                  -> Name, Typ, Datum
  2 Das Angebot auf einen Nenner    -> Preis, Bezug, Sofortzahlung, Anteil
  3 K.-o.-Kriterien vor dem LOI     -> Nachweise, nicht verhandelbar
  4 Bewertung (gewichtet)           -> 11 Kriterien in drei Gruppen
  5 Ergebnis                        -> Gesamtbewertung, Rang, Warnhinweis
  6 Wie du das Ergebnis liest
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- CI-Tokens
NAVY      = "0F2744"
ORANGE    = "EE7A1E"
ORANGE_LT = "FDEBD8"
INK       = "1A1A1A"
GREY      = "666666"
GREY_LT   = "8A94A6"
LINE      = "E5E7EB"
PAPER     = "EEF1F5"
INPUT_BG  = "FFF8EF"
INPUT_FG  = "B35C0F"
WHITE     = "FFFFFF"
NAVY_TEXT = "C7D0DE"
RED       = "C0392B"

FONT = "Calibri"

EUR  = '#,##0 "€";[Red]-#,##0 "€"'
EURC = '#,##0 "€";[Red]-#,##0 "€";""'
PCT  = '0 %'
PCTC = '0 %;-0 %;""'
NUM  = '0;;""'
TXT  = '@'


def f(sz=10.5, b=False, color=INK, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


LEFT   = Alignment(horizontal="left",   vertical="center")
LEFTW  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
LABELW = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

KAND = ["D", "E", "F", "G"]          # bis zu vier Kaufinteressenten
GEW  = "C"                           # Gewichtungsspalte

TYPEN = "Privatperson / MBI,Strategischer Käufer,Finanzinvestor,Nachfolger aus dem Umfeld"
BEZUG = "Unternehmenswert (cash and debt free),Auszahlung an dich"
JANEIN = "ja,nein,offen"
SCORE = "0,1,2,3"

# ---------------------------------------------------------------- Kriterien
# (Kriterium, Gewicht)
GRUPPEN = [
    ("A   Seriosität des Käufers", [
        ("Käuferprofil ausgefüllt und aussagekräftig", 0.05),
        ("Stellt konkrete, projektbezogene Fragen statt Standardtexte", 0.05),
        ("Plausibler Fit zwischen seinen Fähigkeiten und deinem Unternehmen", 0.10),
        ("Antwortet verbindlich und zeitnah (nicht länger als 3–4 Tage)", 0.05),
    ]),
    ("B   Das Angebot", [
        ("Höhe des Kaufpreises im Verhältnis zu deiner Bewertungsbandbreite", 0.15),
        ("Anteil der Sofortzahlung am Kaufpreis (siehe Abschnitt 2)", 0.15),
        ("Wenige und klar formulierte Vorbehalte (Prüfung, Zustimmung Dritter)", 0.05),
        ("Realistischer Zeitrahmen bis zum Abschluss (4–6 Wochen nach Prüfung)", 0.05),
    ]),
    ("C   Finanzierung", [
        ("Anteil Eigenkapital an der Finanzierung", 0.15),
        ("Fremdfinanzierung ist für das Unternehmen tragbar", 0.10),
        ("Finanzierungspartner wurde vor dem LOI eingebunden", 0.10),
    ]),
]

KO = [
    "Vertraulichkeitserklärung unterschrieben",
    "Eigenkapitalnachweis vorgelegt",
    "Bank-Term-Sheet vorgelegt (falls Fremdkapital nötig)",
    "Bonitätsauskunft vorgelegt",
    "Referenzen benannt (1–2 Kontakte)",
]

LESEHILFE = [
    ("Das höchste Angebot ist nicht das beste",
     "Zahlungsmodell und Finanzierung entscheiden über die tatsächliche Qualität eines Angebots, "
     "nicht die Zahl auf dem Papier. Ein niedrigeres Angebot mit hoher Sofortzahlung kann besser sein "
     "als ein höheres mit Earn-Out und vielen Vorbehalten."),
    ("Ein offenes K.-o.-Kriterium schlägt jeden Score",
     "Wer Eigenkapitalnachweis, Bank-Term-Sheet oder Vertraulichkeitserklärung nicht vorlegen kann "
     "oder will, geht nicht in den LOI, egal wie gut die Bewertung aussieht."),
    ("Verkäuferdarlehen und Earn-Out sind normal",
     "Gerade bei kleineren Transaktionen sind gestundete Anteile oft die Voraussetzung dafür, dass die "
     "Finanzierung überhaupt steht. Entscheidend ist nicht, ob sie vorkommen, sondern wie klein der "
     "sofort gezahlte Teil dadurch wird."),
    ("Vergleichbar wird es erst über denselben Bezug",
     "Kläre bei jedem Angebot, ob sich die Zahl auf den Unternehmenswert bezieht oder auf das, was nach "
     "Abzug der Schulden bei dir ankommt. Sonst vergleichst du zwei verschiedene Dinge."),
    ("Halte mehrere Kandidaten parallel im Prozess",
     "Wettbewerb verbessert deine Verhandlungsposition und den Preis. Entscheide dich erst mit dem LOI, "
     "denn ab dort verhandelst du in der Regel exklusiv."),
]

# ---------------------------------------------------------------- Beispiel
BEISPIEL = {
    "name":     ["M. Brandt", "Nordbau Gruppe GmbH", "K. Reimers", None],
    "typ":      ["Privatperson / MBI", "Strategischer Käufer", "Privatperson / MBI", None],
    "datum":    ["14.04.2026", "02.05.2026", "21.05.2026", None],
    "preis":    [780000, 820000, 900000, None],
    "bezug":    ["Unternehmenswert (cash and debt free)"] * 3 + [None],
    "sofort":   [550000, 780000, 450000, None],
    "ko":       [["ja", "ja", "ja", "ja", "ja"],
                 ["ja", "ja", "ja", "ja", "ja"],
                 ["ja", "ja", "nein", "ja", "offen"],
                 [None] * 5],
    "scores":   [[3, 3, 3, 2, 2, 2, 2, 3, 2, 2, 3],
                 [3, 3, 2, 3, 3, 3, 2, 3, 3, 3, 3],
                 [2, 2, 2, 1, 3, 1, 1, 1, 1, 1, 1],
                 [None] * 11],
}
LEER = {k: ([None] * 4 if k not in ("ko", "scores") else
            [[None] * (5 if k == "ko" else 11) for _ in range(4)])
        for k in BEISPIEL}


class Sheet:
    """Schreibt die Matrix zeilenweise und merkt sich die Zeilennummern."""

    def __init__(self, ws, data, beispiel=False):
        self.ws = ws
        self.d = data
        self.beispiel = beispiel
        self.r = 1
        self.ref = {}
        self.dv = {}

    # -- Bausteine ------------------------------------------------
    def row_h(self, h):
        self.ws.row_dimensions[self.r].height = h

    def band(self, h):
        for col in "ABCDEFGH":
            self.ws[f"{col}{self.r}"].fill = fill(NAVY)
        self.row_h(h)
        self.r += 1

    def section(self, nr, titel):
        ws = self.ws
        self.row_h(7)
        self.r += 1
        ws.merge_cells(f"B{self.r}:G{self.r}")
        c = ws[f"B{self.r}"]
        c.value = f"  {nr}   {titel.upper()}"
        c.font = f(11, True, WHITE)
        c.alignment = LEFT
        for col in "BCDEFG":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(24)
        self.r += 1

    def note(self, text, color=GREY, size=9.5, italic=False):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:G{self.r}")
        c = ws[f"B{self.r}"]
        c.value = text
        c.font = f(size, False, color, italic)
        c.alignment = LEFTW
        self.row_h(14)   # AutoFit vergroessert bei Bedarf
        self.r += 1

    def gruppe(self, titel):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:G{self.r}")
        c = ws[f"B{self.r}"]
        c.value = titel
        c.font = f(10, True, NAVY)
        c.alignment = LEFT
        for col in "BCDEFG":
            ws[f"{col}{self.r}"].fill = fill(ORANGE_LT)
        self.row_h(19)
        self.r += 1

    def input_row(self, label, values, fmt=TXT, dv=None, weight=None, indent=False,
                  height=18):
        """Beschriftung ueber B:C (oder B mit Gewicht in C), Eingaben in D:G."""
        ws = self.ws
        if weight is None:
            ws.merge_cells(f"B{self.r}:C{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("      " if indent else "") + label
        c.font = f(10.5)
        c.alignment = LABELW
        c.border = Border(bottom=side(LINE))
        ws[f"C{self.r}"].border = Border(bottom=side(LINE))
        if weight is not None:
            w = ws[f"{GEW}{self.r}"]
            w.value = weight
            w.number_format = PCT
            w.font = f(10.5, True, INPUT_FG)
            w.fill = fill(INPUT_BG)
            w.alignment = CENTER
            w.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        for i, col in enumerate(KAND):
            cell = ws[f"{col}{self.r}"]
            cell.value = values[i] if values else None
            cell.number_format = fmt
            cell.font = f(10.5, False, INPUT_FG)
            cell.fill = fill(INPUT_BG)
            cell.alignment = CENTER if fmt in (TXT, NUM) else RIGHT
            cell.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        if dv:
            self.add_dv(dv, f"{KAND[0]}{self.r}:{KAND[-1]}{self.r}")
        # Zeilen, in die Text eingetippt oder ausgewaehlt wird, brauchen von
        # vornherein Platz fuer zwei Zeilen: der AutoFit-Lauf beim Bauen kann
        # sie nicht messen, solange die Vorlage noch leer ist.
        self.row_h(height)
        row = self.r
        self.r += 1
        return row

    def calc_row(self, label, formula_tpl, fmt=EURC, accent=False, size=10.5, height=19):
        ws = self.ws
        bg = ORANGE_LT if accent else PAPER
        col_ink = NAVY if accent else INK
        ws.merge_cells(f"B{self.r}:C{self.r}")
        c = ws[f"B{self.r}"]
        c.value = label
        if isinstance(label, str) and label.startswith("="):
            c.data_type = "s"          # sonst liest Excel das Label als Formel
        c.font = f(size, True, col_ink)
        c.alignment = LABELW
        for col in "BC":
            ws[f"{col}{self.r}"].fill = fill(bg)
            ws[f"{col}{self.r}"].border = Border(top=side(NAVY if accent else GREY_LT),
                                                 bottom=side(LINE))
        for col in KAND:
            cell = ws[f"{col}{self.r}"]
            cell.value = formula_tpl.format(c=col)
            cell.number_format = fmt
            cell.font = f(size, True, col_ink)
            cell.fill = fill(bg)
            cell.alignment = CENTER
            cell.border = Border(top=side(NAVY if accent else GREY_LT), bottom=side(LINE))
        self.row_h(height)
        row = self.r
        self.r += 1
        return row

    def add_dv(self, liste, bereich):
        if liste not in self.dv:
            d = DataValidation(type="list", formula1=f'"{liste}"', allow_blank=True)
            self.ws.add_data_validation(d)
            self.dv[liste] = d
        self.dv[liste].add(bereich)

    def spacer(self, h=8):
        self.row_h(h)
        self.r += 1

    def kopfzeile_kandidaten(self, label="Kriterium"):
        """Spaltenkopf, der die eingetragenen Namen spiegelt."""
        ws = self.ws
        ws.merge_cells(f"B{self.r}:C{self.r}")
        ws[f"B{self.r}"].value = label
        ws[f"B{self.r}"].font = f(9.5, True, GREY)
        ws[f"B{self.r}"].alignment = Alignment(horizontal="left", vertical="bottom")
        for i, col in enumerate(KAND):
            c = ws[f"{col}{self.r}"]
            c.value = '=IF({c}{n}="","Kandidat {i}",{c}{n})'.format(c=col, n=self.ref["name"], i=i + 1)
            c.font = f(9.5, True, NAVY)
            c.alignment = CENTER
        self.row_h(30)   # Platz fuer zweizeilige Kandidatennamen
        self.r += 1

    # -- Abschnitte -----------------------------------------------
    def kopf(self):
        ws = self.ws
        self.band(6)
        ws[f"B{self.r}"] = "DUB  ·  VERKÄUFER-ACADEMY"
        ws[f"B{self.r}"].font = f(10, True, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        ws.merge_cells(f"E{self.r}:G{self.r}")
        ws[f"E{self.r}"] = "MODUL 6 · KÄUFER- UND ANGEBOTS-SCREENING"
        ws[f"E{self.r}"].font = f(10, True, ORANGE)
        ws[f"E{self.r}"].alignment = RIGHT
        self.band(18)

        ws.merge_cells(f"B{self.r}:G{self.r}")
        ws[f"B{self.r}"] = "Käufer-Bewertungsmatrix" + ("  –  ausgefülltes Beispiel" if self.beispiel else "")
        ws[f"B{self.r}"].font = f(20, True, WHITE)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(30)

        ws.merge_cells(f"B{self.r}:G{self.r}")
        ws[f"B{self.r}"] = "Mehrere Kaufinteressenten systematisch vergleichen, bevor du dich für einen LOI entscheidest"
        ws[f"B{self.r}"].font = f(11, False, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(19)
        self.band(8)
        self.spacer(10)

    def anleitung(self):
        ws = self.ws
        if self.beispiel:
            ws.merge_cells(f"B{self.r}:G{self.r}")
            c = ws[f"B{self.r}"]
            c.value = ("Dieses Blatt zeigt die Matrix mit drei Beispiel-Kandidaten. Beachte den Vergleich von "
                       "Nordbau und K. Reimers: Das höhere Angebot schneidet schlechter ab, weil Sofortzahlung "
                       "und Finanzierung schwächer sind. Arbeite im Blatt „Matrix“ mit deinen eigenen Kandidaten.")
            c.font = f(10, False, NAVY)
            c.alignment = LEFTW
            for col in "BCDEFG":
                ws[f"{col}{self.r}"].fill = fill(ORANGE_LT)
            self.row_h(30)
            self.r += 1
            self.spacer(10)
            return

        self.section("▶", "So gehst du vor")
        for t in [
            "1.  Trage in Abschnitt 1 und 2 ein, wer sich gemeldet hat und was geboten wird. Der Anteil der "
            "Sofortzahlung rechnet sich von selbst.",
            "2.  Hake in Abschnitt 3 die Nachweise ab. Diese Punkte sind nicht verhandelbar: Wer sie nicht "
            "vorlegen kann oder will, kommt nicht in den LOI, unabhängig von seiner Bewertung.",
            "3.  Vergib in Abschnitt 4 für jedes Kriterium 0 bis 3 Punkte je Kandidat. Die Gewichtung ist "
            "vorbelegt und anpassbar, in Summe muss sie 100 % ergeben.",
            "4.  Abschnitt 5 zeigt Gesamtbewertung und Rangfolge, Abschnitt 6 hilft dir beim Einordnen.",
        ]:
            self.note(t, INK, 10)
        self.note("Punkteskala:   0 = nicht erfüllt oder nicht belegt   ·   1 = ansatzweise   ·   "
                  "2 = weitgehend   ·   3 = vollständig erfüllt", INK, 10)
        self.note("Die orange hinterlegten Felder sind Eingabefelder. Alle übrigen Werte berechnen sich automatisch.",
                  GREY_LT, 9.5, italic=True)
        self.spacer(10)

    def abschnitt_1(self):
        ws = self.ws
        self.section("1", "Die Kandidaten")
        heads = ["Kandidat 1", "Kandidat 2", "Kandidat 3", "Kandidat 4"]
        ws.merge_cells(f"B{self.r}:C{self.r}")
        ws[f"B{self.r}"].value = "Angaben aus Anfrage und Käuferprofil"
        ws[f"B{self.r}"].font = f(9.5, True, GREY)
        ws[f"B{self.r}"].alignment = Alignment(horizontal="left", vertical="bottom")
        for col, h in zip(KAND, heads):
            c = ws[f"{col}{self.r}"]
            c.value = h
            c.font = f(9, True, NAVY)
            c.alignment = CENTER
        self.row_h(20)
        self.r += 1

        self.ref["name"] = self.input_row("Name oder Firma", self.d["name"], height=30)
        self.input_row("Käufertyp", self.d["typ"], dv=TYPEN, height=30)
        self.input_row("Datum der Anfrage", self.d["datum"])
        self.note("Bei kleinen Unternehmen sind Privatpersonen und MBI-Kandidaten die häufigsten Käufer. "
                  "Der Typ ist kein Qualitätsmerkmal, er hilft dir nur einzuordnen, was dein Gegenüber vorhat.",
                  GREY_LT, 9.5, italic=True)
        self.spacer(10)

    def abschnitt_2(self):
        self.section("2", "Das Angebot auf einen Nenner")
        self.note("Kläre bei jedem Angebot, worauf sich die Zahl bezieht. Ein Angebot über den Unternehmenswert "
                  "klingt höher, obwohl nach Abzug der Schulden weniger bei dir ankommt (siehe Modul 4, Preisbrücke).",
                  INK, 10)
        self.kopfzeile_kandidaten("Angaben zum Angebot")
        self.input_row("Gebotener Kaufpreis", self.d["preis"], fmt=EUR)
        self.ref["preis"] = self.r - 1
        self.input_row("Bezug der Zahl", self.d["bezug"], dv=BEZUG, height=30)
        self.input_row("Davon Sofortzahlung bei Closing", self.d["sofort"], fmt=EUR)
        self.ref["sofort"] = self.r - 1
        self.calc_row("= Anteil der Sofortzahlung am Kaufpreis",
                      '=IFERROR({{c}}{s}/{{c}}{p},"")'.format(s=self.ref["sofort"], p=self.ref["preis"]),
                      fmt=PCTC, accent=True, size=11, height=21)
        self.note("Der Rest ist gestundet, meist als Verkäuferdarlehen oder Earn-Out. Beides ist normal und oft "
                  "nötig, damit die Finanzierung steht. Kritisch wird es, wenn der sofort gezahlte Teil so klein "
                  "wird, dass du den Großteil des Risikos trägst.", GREY_LT, 9.5, italic=True)
        self.spacer(10)

    def abschnitt_3(self):
        self.section("3", "K.-o.-Kriterien vor dem LOI")
        self.note("Diese Nachweise sind nicht verhandelbar. Fordere sie spätestens an, bevor du dich für einen "
                  "Käufer entscheidest und damit in der Regel exklusiv verhandelst.", INK, 10)
        self.kopfzeile_kandidaten("Nachweis")
        start = self.r
        for i, label in enumerate(KO):
            self.input_row(label, [self.d["ko"][k][i] for k in range(4)],
                           dv=JANEIN, indent=True)
        ende = self.r - 1
        self.ref["ko"] = (start, ende)
        self.calc_row("= Status der Nachweise",
                      '=IF(COUNTA({{c}}{a}:{{c}}{b})=0,"",'
                      'IF(COUNTIF({{c}}{a}:{{c}}{b},"nein")>0,'
                      'COUNTIF({{c}}{a}:{{c}}{b},"nein")&"× nein",'
                      'IF(COUNTIF({{c}}{a}:{{c}}{b},"ja")={n},"alle erfüllt",'
                      'COUNTIF({{c}}{a}:{{c}}{b},"offen")&"× offen")))'.format(a=start, b=ende, n=len(KO)),
                      fmt=TXT, accent=True, size=11, height=21)
        self.ref["ko_status"] = self.r - 1
        self.spacer(10)

    def abschnitt_4(self):
        ws = self.ws
        self.section("4", "Bewertung der Kandidaten (gewichtet)")
        self.note("Vergib je Kriterium 0 bis 3 Punkte:   0 = nicht erfüllt oder nicht belegt   ·   "
                  "1 = ansatzweise   ·   2 = weitgehend   ·   3 = vollständig erfüllt", INK, 10)

        # Kopfzeile mit Gewichtungsspalte
        ws.merge_cells(f"B{self.r}:B{self.r}")
        ws[f"B{self.r}"].value = "Kriterium"
        ws[f"B{self.r}"].font = f(9.5, True, GREY)
        ws[f"B{self.r}"].alignment = Alignment(horizontal="left", vertical="bottom")
        ws[f"{GEW}{self.r}"].value = "Gewicht"
        ws[f"{GEW}{self.r}"].font = f(9, True, NAVY)
        ws[f"{GEW}{self.r}"].alignment = CENTER
        for i, col in enumerate(KAND):
            c = ws[f"{col}{self.r}"]
            c.value = '=IF({c}{n}="","Kandidat {i}",{c}{n})'.format(c=col, n=self.ref["name"], i=i + 1)
            c.font = f(9.5, True, NAVY)
            c.alignment = CENTER
        self.row_h(30)   # Platz fuer zweizeilige Kandidatennamen
        self.r += 1

        rows = []
        idx = 0
        for titel, kriterien in GRUPPEN:
            self.gruppe(titel)
            for label, gewicht in kriterien:
                werte = [self.d["scores"][k][idx] for k in range(4)]
                rows.append(self.input_row(label, werte, fmt=NUM, dv=SCORE,
                                           weight=gewicht, indent=True))
                idx += 1
        self.ref["krit"] = (rows[0], rows[-1])
        self.ref["n_krit"] = len(rows)

        a, b = self.ref["krit"]
        # Gewichtungssumme pruefen
        ws.merge_cells(f"B{self.r}:C{self.r}")
        chk = ws[f"B{self.r}"]
        chk.value = ('=IF(SUM({g}{a}:{g}{b})=1,"Gewichtung ergibt 100 % – passt.",'
                     '"Achtung: Gewichtung ergibt "&TEXT(SUM({g}{a}:{g}{b}),"0 %")&". Bitte anpassen.")'
                     ).format(g=GEW, a=a, b=b)
        chk.font = f(9.5, False, GREY_LT, True)
        chk.alignment = LEFTW
        self.row_h(15)
        self.r += 1
        self.spacer(6)

        self.calc_row("Bewertete Kriterien",
                      '=IF(COUNT({{c}}{a}:{{c}}{b})=0,"",COUNT({{c}}{a}:{{c}}{b})&" von {n}")'.format(
                          a=a, b=b, n=self.ref["n_krit"]),
                      fmt=TXT, size=10, height=18)
        self.spacer(10)

    def abschnitt_5(self):
        ws = self.ws
        a, b = self.ref["krit"]
        self.section("5", "Ergebnis")
        self.kopfzeile_kandidaten("Auswertung")

        # Gesamtbewertung: gewichtete Punkte, auf 100 % normiert (max 3 Punkte)
        ws.merge_cells(f"B{self.r}:C{self.r}")
        c = ws[f"B{self.r}"]
        c.value = "GESAMTBEWERTUNG"
        c.font = f(13, True, WHITE)
        c.alignment = LEFT
        for col in KAND:
            cell = ws[f"{col}{self.r}"]
            cell.value = ('=IF(COUNT({{c}}{a}:{{c}}{b})=0,"",'
                          'SUMPRODUCT({{c}}{a}:{{c}}{b},${g}${a}:${g}${b})/3)').format(
                              a=a, b=b, g=GEW).format(c=col)
            cell.number_format = PCTC
            cell.font = f(13, True, WHITE)
            cell.alignment = CENTER
        for col in "BCDEFG":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(32)
        tot = self.r
        self.r += 1
        self.ref["total"] = tot

        self.calc_row("Rang",
                      '=IF({{c}}{t}="","",RANK({{c}}{t},${d}${t}:${g}${t}))'.format(
                          t=tot, d=KAND[0], g=KAND[-1]),
                      fmt=TXT, size=11, height=19)

        # Warnzeile
        ws.merge_cells(f"B{self.r}:C{self.r}")
        c = ws[f"B{self.r}"]
        c.value = "Hinweis"
        c.font = f(10, True, NAVY)
        c.alignment = LABELW
        for col in "BC":
            ws[f"{col}{self.r}"].fill = fill(PAPER)
        for col in KAND:
            cell = ws[f"{col}{self.r}"]
            cell.value = ('=IF({c}{t}="","",IF(ISNUMBER(SEARCH("nein",{c}{k})),'
                          '"nicht für den LOI",IF(ISNUMBER(SEARCH("offen",{c}{k})),'
                          '"Nachweise offen","")))').format(c=col, t=tot, k=self.ref["ko_status"])
            cell.number_format = TXT
            cell.font = f(10, True, RED)
            cell.fill = fill(PAPER)
            cell.alignment = CENTER
        self.row_h(19)
        self.r += 1

        self.note("Ein Kandidat mit einem offenen oder verneinten K.-o.-Kriterium kommt nicht in den LOI, "
                  "unabhängig von seiner Gesamtbewertung.", GREY_LT, 9.5, italic=True)
        self.spacer(10)

    def abschnitt_6(self):
        ws = self.ws
        self.section("6", "Wie du das Ergebnis liest")
        for titel, text in LESEHILFE:
            b = ws[f"B{self.r}"]
            b.value = titel
            b.font = f(10, True, NAVY)
            b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            b.border = Border(top=side(LINE))
            ws.merge_cells(f"C{self.r}:G{self.r}")
            t = ws[f"C{self.r}"]
            t.value = text
            t.font = f(9.5, False, INK)
            t.alignment = LEFTW
            for col in "CDEFG":
                ws[f"{col}{self.r}"].border = Border(top=side(LINE))
            self.row_h(30)
            self.r += 1
        self.spacer(12)

    def fuss(self):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:G{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("Diese Matrix dient der Orientierung und ersetzt keine Rechts-, Steuer- oder Finanzberatung.    "
                   "DUB · info@dub.de · www.dub.de    © 2026 Deutsche Unternehmerbörse DUB.de GmbH")
        c.font = f(9, False, GREY_LT)
        c.alignment = LEFTW
        for col in "BCDEFG":
            ws[f"{col}{self.r}"].border = Border(top=side(LINE))
        self.row_h(24)
        self.r += 1

    def layout(self):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = NAVY
        for col, w in {"A": 2.2, "B": 58, "C": 13,
                       "D": 17, "E": 17, "F": 17, "G": 17, "H": 2.2}.items():
            ws.column_dimensions[col].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4

    def build(self):
        self.layout()
        self.kopf()
        self.anleitung()
        self.abschnitt_1()
        self.abschnitt_2()
        self.abschnitt_3()
        self.abschnitt_4()
        self.abschnitt_5()
        self.abschnitt_6()
        self.fuss()


# ------------------------------------------------------------------ AutoFit
AUTOFIT_PS = r"""
param([string]$Path)
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
$wb = $xl.Workbooks.Open($Path)
$HELP = 26
foreach ($ws in $wb.Worksheets) {
    $used = $ws.UsedRange
    $last = $used.Row + $used.Rows.Count - 1
    $minH = @{}
    for ($r = 1; $r -le $last; $r++) { $minH[$r] = $ws.Rows.Item($r).RowHeight }
    for ($r = 1; $r -le $last; $r++) {
        $text = $null; $fromCol = 0; $span = 0
        for ($c = 2; $c -le 7; $c++) {
            $cell = $ws.Cells.Item($r, $c)
            if ($cell.MergeCells -and $cell.WrapText) {
                $ma = $cell.MergeArea
                if ($ma.Rows.Count -eq 1 -and $ma.Row -eq $r) {
                    $tl = $ws.Cells.Item($r, $ma.Column)
                    if ($tl.Value2 -is [string] -and $tl.Value2.Length -gt 0) {
                        $text = $tl.Value2; $fromCol = $ma.Column; $span = $ma.Columns.Count
                    }
                }
                break
            }
        }
        if ($null -ne $text) {
            $w = 0
            for ($k = 0; $k -lt $span; $k++) { $w += $ws.Columns.Item($fromCol + $k).ColumnWidth }
            $h = $ws.Cells.Item($r, $fromCol)
            $ws.Columns.Item($HELP).ColumnWidth = $w
            $probe = $ws.Cells.Item($r, $HELP)
            $probe.Font.Name = $h.Font.Name; $probe.Font.Size = $h.Font.Size
            $probe.Font.Bold = $h.Font.Bold; $probe.WrapText = $true
            $probe.Value2 = $text
        }
        $fit = ($null -ne $text)
        if (-not $fit) {
            for ($c = 2; $c -le 7; $c++) {
                $cell = $ws.Cells.Item($r, $c)
                if ((-not $cell.MergeCells) -and $cell.WrapText -and $null -ne $cell.Value2) { $fit = $true; break }
            }
        }
        if ($fit) {
            $ws.Rows.Item($r).AutoFit() | Out-Null
            $ws.Rows.Item($r).RowHeight = [Math]::Max($ws.Rows.Item($r).RowHeight + 3, $minH[$r])
            if ($null -ne $text) { $ws.Cells.Item($r, $HELP).Clear() | Out-Null }
        }
    }
    $ws.Columns.Item($HELP).ColumnWidth = 8.43
}
$wb.Worksheets.Item(1).Activate()
$wb.Save(); $wb.Close($true); $xl.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null
"""


def autofit(path):
    """Zeilenhoehen von Excel messen lassen (nur Windows mit installiertem Excel)."""
    import subprocess, tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as fh:
        fh.write(AUTOFIT_PS)
        script = fh.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", script, "-Path", os.path.abspath(path)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            print("AutoFit uebersprungen (kein Excel?):", (r.stderr or "").strip()[:200])
        else:
            print("Zeilenhoehen von Excel gemessen.")
    finally:
        os.unlink(script)


def main():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Matrix"
    Sheet(ws1, LEER, beispiel=False).build()

    ws2 = wb.create_sheet("Beispiel")
    Sheet(ws2, BEISPIEL, beispiel=True).build()
    ws2.sheet_properties.tabColor = ORANGE

    wb.properties.title = "DUB Käufer-Bewertungsmatrix"
    wb.properties.creator = "Deutsche Unternehmerbörse DUB.de GmbH"
    wb.properties.description = "Verkäufer-Academy, Modul 6: Käufer- und Angebots-Screening"

    out = os.path.join("Downloads", "Verkaeufer", "kaeufer-bewertungsmatrix.xlsx")
    wb.save(out)
    print("geschrieben:", out)
    autofit(out)


if __name__ == "__main__":
    main()
