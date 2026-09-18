# -*- coding: utf-8 -*-
"""
DUB Kaeufer-Academy - DD-Checkliste (Modul 6)
Erzeugt Downloads/Kaeufer/dd-checkliste.xlsx

Zwei Blaetter:
  "Kernpruefung"   - Financial, Tax, Legal. Gilt fuer jeden Deal, Schwerpunkt
                     bis etwa 2 Mio. EUR. Pendant zum Verkaeufer-Leitfaden
                     "So laeuft eine Kaeuferpruefung (Due Diligence)".
  "Zusatzbausteine" - optionale Pruefbereiche je nach Transaktion, aus dem
                     DealCircle-Guide "Due Diligence Checklist", ergaenzt um
                     eine KI-Due-Diligence.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

# ---------------------------------------------------------------- CI-Tokens
NAVY      = "0F2744"
ORANGE    = "EE7A1E"
ORANGE_LT = "FDEBD8"
INK       = "1A1A1A"
GREY      = "666666"
GREY_LT   = "8A94A6"
LINE      = "E5E7EB"
INPUT_BG  = "FFF8EF"
INPUT_FG  = "B35C0F"
WHITE     = "FFFFFF"
NAVY_TEXT = "C7D0DE"
FONT = "Calibri"


def f(sz=10.5, b=False, color=INK, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
LEFTW  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

STATUS = '"offen,in Prüfung,geklärt,Klärung nötig"'


# ---------------------------------------------------------------- Kernprüfung
FINANCIAL = [
    ("Bereinigtes EBITDA",
     "Lass dir jede Bereinigung belegen. Größter Posten ist meist das Geschäftsführergehalt, "
     "normalisiert auf marktübliches Niveau. Was nicht belegt ist, gehört nicht in deine Bewertung."),
    ("Konsistenz der Zahlen",
     "Jahresabschlüsse, laufende Auswertungen (BWA), Summen- und Saldenliste und Steuererklärungen "
     "müssen zueinander passen. Abweichungen erklären lassen, auch harmlose."),
    ("Entwicklung im laufenden Jahr",
     "Passt die aktuelle BWA zum letzten Jahresabschluss? Umsatzsprünge und Margeneinbrüche brauchen "
     "eine nachvollziehbare Begründung."),
    ("Offene Forderungen",
     "Viele überfällige Kundenrechnungen deuten auf Zahlungsprobleme der Kunden hin oder darauf, "
     "dass zu spät gemahnt wird. Beides übernimmst du mit."),
    ("Im Tagesgeschäft gebundenes Kapital",
     "Wie viel Geld steckt normalerweise in Vorräten und offenen Rechnungen? Dieser Normalwert geht in "
     "die Preisbrücke ein (Fachbegriff: Working Capital)."),
    ("Finanzverbindlichkeiten und Sicherheiten",
     "Darlehen samt Bürgschaften, Grundschulden und Sicherungsübereignungen. Sie mindern den "
     "Kaufpreis über die Preisbrücke."),
    ("Nicht bilanzierte Verpflichtungen",
     "Bürgschaften, Pensionszusagen, Leasing. Sie stehen nicht in der Bilanz und treffen dich trotzdem. "
     "Frage ausdrücklich danach."),
    ("Investitionsstau",
     "Anlagenverzeichnis mit Alter und Restbuchwert, dazu die Investitionen der letzten Jahre. Was du "
     "bald ersetzen musst, gehört in die Preisbrücke."),
]

TAX = [
    ("Offene oder angekündigte Betriebsprüfung",
     "Stand der Prüfung und mögliches Nachzahlungsrisiko. Beim Share Deal übernimmst du die "
     "Steuerhistorie mit, beim Asset Deal in der Regel nicht."),
    ("Verträge mit nahestehenden Personen",
     "Bei inhabergeführten Firmen der wichtigste Prüfpunkt: Ehepartner auf der Lohnliste, Miete an die "
     "private Immobilie, private Fahrzeuge. Steuerlich eine verdeckte Gewinnausschüttung."),
    ("Umsatzsteuer",
     "Falsch behandelte Rechnungen und Sonderfälle summieren sich über Jahre. Was nicht korrigiert ist, "
     "landet als Risiko im Kaufvertrag."),
    ("Steuerbescheide und Prüfberichte",
     "Die letzten Bescheide und die Berichte abgeschlossener Prüfungen zeigen, ob eine Nachzahlung "
     "im Raum steht."),
    ("Lohnsteuer- und Sozialversicherungsprüfung",
     "Läuft turnusmäßig und ist bei kleinen Betrieben die häufigste Quelle für Nachzahlungen. "
     "Lass dir die Berichte der letzten Prüfungen zeigen."),
    ("Fördermittel und Zuschüsse",
     "Bindungsfristen und Verwendungsauflagen prüfen. Manche Zuschüsse sind zurückzuzahlen, wenn der "
     "Betrieb den Eigentümer wechselt."),
]

LEGAL = [
    ("Change-of-Control-Klauseln",
     "Kunden-, Liefer-, Miet- und Leasingverträge, die dem Vertragspartner ein Kündigungsrecht geben, "
     "sobald die Firma den Eigentümer wechselt. Der wichtigste Prüfpunkt überhaupt."),
    ("Lückenlose Eigentümerkette",
     "Alle früheren Anteilsübertragungen müssen formgerecht dokumentiert sein. Ein Formfehler, etwa "
     "eine fehlende notarielle Beurkundung, macht die Übertragung unwirksam: Deal Breaker."),
    ("Genehmigungen und Konzessionen",
     "Manche behördlichen Erlaubnisse hängen an der Person oder Qualifikation des Inhabers "
     "(z.B. Handwerksrolle, Gaststättenerlaubnis, Paragraf 34c GewO) und gehen nicht automatisch über."),
    ("Geistiges Eigentum",
     "Marken, Domains, Patente und Lizenzen müssen dem Unternehmen gehören, nicht dem Inhaber privat. "
     "Sonst kaufst du eine Firma ohne ihre Marke."),
    ("Laufende Rechtsstreitigkeiten",
     "Auch angekündigte gehören auf den Tisch. Bei existenzbedrohendem Streitwert ist das kein "
     "Preisthema, sondern ein Abbruchgrund."),
    ("Arbeitsverhältnisse",
     "Kündigungsfristen, Sonderzusagen und die Frage, ob Schlüsselpersonen gebunden sind. Du willst "
     "wissen, welche Leute und welche Kosten du übernimmst."),
    ("Gesellschaftsvertrag und Gesellschafterliste",
     "Ist der Verkäufer verkaufsberechtigt? Bei mehreren Gesellschaftern: sind alle dabei? Gibt es "
     "Vorkaufsrechte oder Zustimmungsvorbehalte?"),
]

AUSWERTUNG = [
    ("Befunde einordnen",
     "Red Flag oder Deal Breaker? Ein echter Deal Breaker, etwa systematische Bilanzverschleierung, ist "
     "kein Anlass für eine Preisreduktion, sondern ein Grund abzubrechen."),
    ("Bewertung nachziehen",
     "Arbeite die Befunde in deine Preisbrücke ein. Nur bei echten, substanziellen Befunden: Wer den "
     "Preis ohne triftigen Grund aufmacht, verliert das Vertrauen der Gegenseite."),
    ("Risiken in den Kaufvertrag",
     "Was du gefunden hast und nicht über den Preis regelst, gehört in Garantien, Freistellungen oder "
     "einen Kaufpreiseinbehalt."),
    ("Finanzierung gegenprüfen",
     "Ändert sich der Kaufpreis, ändert sich dein Finanzierungsbedarf. Sprich mit deiner Bank, bevor "
     "du den Kaufvertrag verhandelst."),
]

KERN = [
    ("1", "Financial DD: die Zahlen", None, FINANCIAL),
    ("2", "Tax DD: die Steuern", None, TAX),
    ("3", "Legal DD: das Recht", None, LEGAL),
    ("4", "Auswertung und nächste Schritte", None, AUSWERTUNG),
]


# ---------------------------------------------------------------- Zusatzbausteine
COMMERCIAL = [
    ("Belastbarkeit der Kundenbeziehungen",
     "Wie fest sind die größten Kunden gebunden, und was passiert nach dem Eigentümerwechsel? Der "
     "Verlust eines wesentlichen Kunden nach der Übernahme ist das größte kundenseitige Risiko."),
    ("Annahmen hinter der Planung",
     "Worauf beruhen die geplanten Zahlen: mehr Kunden, höhere Preise, neue Produkte? Jede Annahme "
     "einzeln prüfen, statt die Summe zu glauben."),
    ("Markt und Wettbewerb",
     "Wächst der Markt oder das Unternehmen? Wer sind die Wettbewerber, und was hält sie fern? "
     "Branchenverbände sind eine gute und kostenlose Quelle."),
]

IT_TECH = [
    ("Abhängigkeit vom System",
     "Was kostet ein Ausfall von einem Tag? Je stärker das Geschäft an der Technik hängt, desto eher "
     "brauchst du hier einen Spezialisten statt dein eigenes Urteil."),
    ("Eigentum an Software und Daten",
     "Gehören Quellcode, Shop-System, Kundendatenbank und Domains dem Unternehmen? Bei extern "
     "entwickelter Software: Sind die Nutzungsrechte übertragbar?"),
    ("Wartungsstand und Sicherheit",
     "Aktuelle Systeme, Datensicherung, dokumentierte Zugänge. Ein Investitionsstau in der Technik "
     "ist so real wie einer bei den Maschinen."),
]

KI = [
    ("Ersetzbarkeit der Leistung",
     "Kann KI das, wofür Kunden heute zahlen, in wenigen Jahren deutlich billiger? Besonders bei "
     "Text-, Übersetzungs-, Gestaltungs- und einfachen Auswertungsleistungen. Du kaufst künftige "
     "Erträge, nicht vergangene."),
    ("Eingesetzte Werkzeuge",
     "Welche KI-Dienste laufen im Betrieb, wofür, und auf welchen Verträgen? Häufig laufen sie auf "
     "Privatkonten einzelner Mitarbeiter statt auf Firmenkonten und sind nach der Übergabe weg."),
    ("Vertrauliche Daten in offenen Diensten",
     "Wurden Kunden-, Mitarbeiter- oder Vertragsdaten in frei zugängliche KI-Dienste eingegeben? Ein "
     "Datenschutzverstoß aus der Vergangenheit geht beim Share Deal auf dich über."),
    ("Rechte an erzeugten Inhalten",
     "Bei KI-erzeugtem Text, Bild oder Code ist der urheberrechtliche Schutz unklar. Prüfe, ob "
     "wesentliche Inhalte des Unternehmens darauf beruhen."),
    ("Abhängigkeit von einem Anbieter",
     "Hängt ein Arbeitsschritt an einem einzelnen Dienst? Preis, Verfügbarkeit und Bedingungen kann "
     "der Anbieter jederzeit ändern, und du hast keinen Vertrag darüber."),
]

HR_ORG = [
    ("Schlüsselpersonen",
     "Wer trägt das Geschäft außer dem Inhaber, und was hält diese Leute? Ohne Bindung übernimmst du "
     "eine Firma, die in Monaten eine andere ist."),
    ("Verdeckte Personalkosten",
     "Sonderzusagen, Boni, Altersvorsorge, Überstundenkonten und Urlaubsrückstände. Vieles davon steht "
     "nicht im Arbeitsvertrag und nicht in der Bilanz."),
    ("Übergang der Arbeitsverhältnisse",
     "Beim Asset Deal gehen sie nach Paragraf 613a BGB mit über und die Mitarbeiter können "
     "widersprechen. Kläre vorher, wen du wirklich brauchst."),
]

UMWELT_IMMOBILIEN = [
    ("Altlasten",
     "Bei Produktion, Werkstatt, Lager oder Tankanlagen: Gibt es Bodenbelastungen aus der "
     "Vergangenheit? Die Sanierungspflicht trifft den Eigentümer, unabhängig vom Verursacher."),
    ("Zustand und Rechte an der Immobilie",
     "Bei eigener Immobilie: Grundbuch, Baulasten, Genehmigungen und Instandhaltungsstau. Bei Miete: "
     "Laufzeit, Kündigungsrechte und ob der Vertrag den Eigentümerwechsel überlebt."),
]

ESG = [
    ("Anforderungen aus der Lieferkette",
     "Verlangen große Kunden Nachweise zu Umwelt, Arbeitsbedingungen oder Herkunft? Fehlen sie, kann "
     "das Aufträge kosten, auch wenn heute noch niemand fragt."),
    ("Anforderungen der Bank",
     "Finanzierer fragen zunehmend nach Umwelt- und Sozialkriterien. Kläre früh, ob deine Finanzierung "
     "davon abhängt."),
]

OPERATIONAL = [
    ("Abläufe und dokumentiertes Wissen",
     "Wie läuft die Leistungserbringung wirklich? Was steht dokumentiert, was nur in Köpfen? "
     "Undokumentierte Abläufe sind der häufigste Grund, warum Übernahmen im ersten Jahr holpern."),
    ("Zustand der Betriebsmittel",
     "Maschinen, Fahrzeuge, Werkzeuge vor Ort ansehen, nicht nur im Anlagenverzeichnis. Wartungsstand, "
     "Ersatzteilversorgung und Alter sagen mehr als der Restbuchwert."),
    ("Lieferanten und Einkauf",
     "Hängt die Leistung an einzelnen Lieferanten? Gibt es Rahmenverträge oder nur Zuruf? Welche "
     "Konditionen gelten nach dem Eigentümerwechsel weiter?"),
    ("Kapazität und Auslastung",
     "Wie viel mehr ginge mit dem, was vorhanden ist? Freie Kapazität ist der günstigste "
     "Wachstumshebel, fehlende bedeutet Investition kurz nach dem Kauf."),
]

VERSICHERUNG = [
    ("Bestehender Schutz",
     "Welche Policen laufen, mit welchen Summen und Selbstbehalten? Betriebshaftpflicht, Inhalt und "
     "Betriebsunterbrechung, je nach Branche Berufshaftpflicht oder Umwelthaftung."),
    ("Deckungslücken",
     "Gibt es Risiken ohne Versicherung? Bei kleinen Betrieben häufig Betriebsunterbrechung, "
     "Produkthaftung und Cyber. Was nicht gedeckt ist, trägst nach der Übernahme du."),
    ("Übergang der Verträge",
     "Beim Share Deal laufen die Policen in der Gesellschaft weiter, beim Asset Deal in der Regel nicht "
     "automatisch. Eine Deckungslücke am Übergabetag kann teuer werden."),
    ("Schadenhistorie",
     "Welche Schäden wurden in den letzten Jahren gemeldet? Häufige Schäden treiben die Prämie und "
     "zeigen operative Schwachstellen, die sonst niemand erwähnt."),
]

ZUSATZ = [
    ("1", "Commercial DD: Geschäftsmodell und Markt",
     "Relevant, wenn du für Wachstum bezahlst oder die Planung des Verkäufers Teil deines Preises ist.",
     COMMERCIAL),
    ("2", "Operational DD: Abläufe und Betriebsmittel",
     "Relevant bei Produktion, Handwerk und Logistik, überall dort, wo die Leistung an Anlagen, "
     "Abläufen und Lieferanten hängt.",
     OPERATIONAL),
    ("3", "IT und Technik",
     "Relevant, wenn die Technik das Geschäft trägt: Onlinehandel, Software, Plattformen, "
     "produzierende Betriebe mit vernetzten Anlagen.",
     IT_TECH),
    ("4", "KI-Prüfung",
     "Relevant für fast jedes Dienstleistungsgeschäft. Dieses Thema steht in keinem klassischen "
     "Due-Diligence-Katalog, entscheidet aber zunehmend über den Wert der Erträge, die du kaufst.",
     KI),
    ("5", "Personal und Organisation",
     "Relevant ab etwa zehn Mitarbeitern oder wenn das Geschäft an wenigen Köpfen hängt.",
     HR_ORG),
    ("6", "Versicherungen",
     "Relevant bei jedem Betrieb mit eigenen Anlagen, Fahrzeugen oder Personal. Schnell geprüft und "
     "deckt Lücken auf, die sonst erst im Schadensfall auffallen.",
     VERSICHERUNG),
    ("7", "Umwelt und Immobilien",
     "Relevant bei eigener Immobilie, bei Produktion und überall dort, wo mit Stoffen gearbeitet wird.",
     UMWELT_IMMOBILIEN),
    ("8", "ESG: Umwelt, Soziales und Unternehmensführung",
     "Relevant, wenn große Kunden oder deine Bank danach fragen. Bei kleinen Betrieben sonst selten "
     "ein eigener Prüfbereich.",
     ESG),
]


class Sheet:
    """Schreibt ein Checklisten-Blatt zeilenweise."""

    def __init__(self, ws, titel, untertitel, abschnitte, anleitung):
        self.ws = ws
        self.titel = titel
        self.untertitel = untertitel
        self.abschnitte = abschnitte
        self.anleitung_zeilen = anleitung
        self.r = 1
        self.breaks = []

    # -- Bausteine ------------------------------------------------
    def row_h(self, h):
        self.ws.row_dimensions[self.r].height = h

    def band(self, h):
        for col in "ABCDEF":
            self.ws[f"{col}{self.r}"].fill = fill(NAVY)
        self.row_h(h)
        self.r += 1

    def section(self, nr, titel):
        ws = self.ws
        self.breaks.append(self.r)
        self.row_h(7)
        self.r += 1
        ws.merge_cells(f"B{self.r}:E{self.r}")
        c = ws[f"B{self.r}"]
        c.value = f"  {nr}   {titel.upper()}"
        c.font = f(11, True, WHITE)
        c.alignment = LEFT
        for col in "BCDE":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(24)
        self.r += 1

    def note(self, text, color=GREY, size=9.5, italic=False, bg=None):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:E{self.r}")
        c = ws[f"B{self.r}"]
        c.value = text
        c.font = f(size, False, color, italic)
        c.alignment = LEFTW
        if bg:
            for col in "BCDE":
                ws[f"{col}{self.r}"].fill = fill(bg)
        self.row_h(14)
        self.r += 1

    def spalten_kopf(self):
        ws = self.ws
        for col, t in zip("BCDE", ["Prüfpunkt", "Worauf du achtest", "Status", "Befund / offene Frage"]):
            c = ws[f"{col}{self.r}"]
            c.value = t
            c.font = f(9.5, True, NAVY)
            c.alignment = Alignment(horizontal="center" if col == "D" else "left",
                                    vertical="bottom", wrap_text=True)
            c.border = Border(bottom=side(ORANGE, "medium"))
        self.row_h(20)
        self.r += 1

    def item(self, titel, text):
        ws = self.ws
        b = ws[f"B{self.r}"]
        b.value = titel
        b.font = f(10.5, True, NAVY)
        b.alignment = LEFTW
        c = ws[f"C{self.r}"]
        c.value = text
        c.font = f(10, False, INK)
        c.alignment = LEFTW
        d = ws[f"D{self.r}"]
        d.value = "offen"
        d.font = f(10, False, INPUT_FG)
        d.fill = fill(INPUT_BG)
        d.alignment = CENTER
        e = ws[f"E{self.r}"]
        e.fill = fill(INPUT_BG)
        e.alignment = LEFTW
        e.font = f(10, False, INPUT_FG)
        for col in "BCDE":
            ws[f"{col}{self.r}"].border = Border(bottom=side(LINE))
        self.row_h(30)
        row = self.r
        self.r += 1
        return row

    def spacer(self, h=8):
        self.row_h(h)
        self.r += 1

    # -- Abschnitte -----------------------------------------------
    def kopf(self):
        ws = self.ws
        self.band(6)
        ws[f"B{self.r}"] = "DUB  ·  KÄUFER-ACADEMY"
        ws[f"B{self.r}"].font = f(10, True, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        ws.merge_cells(f"D{self.r}:E{self.r}")
        ws[f"D{self.r}"] = "MODUL 6 · DUE DILIGENCE MEISTERN"
        ws[f"D{self.r}"].font = f(10, True, ORANGE)
        ws[f"D{self.r}"].alignment = Alignment(horizontal="right", vertical="center")
        self.band(18)

        ws.merge_cells(f"B{self.r}:E{self.r}")
        ws[f"B{self.r}"] = self.titel
        ws[f"B{self.r}"].font = f(20, True, WHITE)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(30)

        ws.merge_cells(f"B{self.r}:E{self.r}")
        ws[f"B{self.r}"] = self.untertitel
        ws[f"B{self.r}"].font = f(11, False, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(19)
        self.band(8)
        self.spacer(10)

    def anleitung(self):
        self.section("▶", "So gehst du vor")
        for t in self.anleitung_zeilen:
            self.note(t, INK, 10)
        self.note("Die orange hinterlegten Felder füllst du aus: Status je Prüfpunkt, daneben dein Befund "
                  "oder die offene Frage an den Verkäufer.", GREY_LT, 9.5, italic=True)
        self.spacer(10)

    def build(self):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = NAVY
        for col, w in {"A": 2.2, "B": 34, "C": 74, "D": 15, "E": 40, "F": 2.2}.items():
            ws.column_dimensions[col].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4

        self.kopf()
        self.anleitung()

        bereiche = []
        for nr, titel, wann, eintraege in self.abschnitte:
            self.section(nr, titel)
            if wann:
                self.note("Wann dieses Thema schwer wiegt: " + wann, NAVY, 10, bg=ORANGE_LT)
                self.spacer(4)
            self.spalten_kopf()
            erste = self.r
            for t, x in eintraege:
                self.item(t, x)
            bereiche.append((erste, self.r - 1))
            self.spacer(10)

        # Fuss
        ws.merge_cells(f"B{self.r}:E{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("Diese Checkliste dient der Orientierung und ersetzt keine Rechts-, Steuer- oder "
                   "Finanzberatung.    DUB · info@dub.de · www.dub.de    "
                   "© 2026 Deutsche Unternehmerbörse DUB.de GmbH")
        c.font = f(9, False, GREY_LT)
        c.alignment = LEFTW
        for col in "BCDE":
            ws[f"{col}{self.r}"].border = Border(top=side(LINE))
        self.row_h(24)
        self.r += 1

        dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
        dv.promptTitle = "Status"
        dv.prompt = "Status wählen"
        ws.add_data_validation(dv)
        for erste, letzte in bereiche:
            dv.add(f"D{erste}:D{letzte}")

        for r in self.breaks[1:]:
            ws.row_breaks.append(Break(id=r))


ANLEITUNG_KERN = [
    "1.  Beauftrage deine Berater, bevor der Datenraum aufgeht. Bei Deals unter 2 Mio. EUR deckt meist "
    "ein Steuerberater Finanzen und Steuern ab und ein Anwalt das Recht.",
    "2.  Fordere die Unterlagen schriftlich an und arbeite den Datenraum Bereich für Bereich durch. "
    "Trage deine Befunde direkt in die letzte Spalte ein.",
    "3.  Bündle deine Rückfragen in einer Liste, statt sie einzeln zu stellen. Das spart beiden Seiten "
    "Zeit und macht den Verlauf nachvollziehbar.",
    "4.  Werte am Ende Abschnitt 4 aus: Was ist Verhandlungsmasse, was ein Abbruchgrund, und was gehört "
    "in den Kaufvertrag?",
    "Prüfe nicht alles, sondern das Richtige: Bei kleinen Deals suchst du Deal Breaker, nicht "
    "Vollständigkeit. Welche Themen darüber hinaus in deinem Fall schwer wiegen, zeigt das zweite Blatt.",
]

ANLEITUNG_ZUSATZ = [
    "Financial, Tax und Legal sind das Pflichtprogramm, sie stehen auf dem ersten Blatt. Darüber hinaus "
    "bringt jedes Unternehmen eigene Chancen und Risiken mit, und die verschieben Schwerpunkt und Umfang "
    "der Prüfung.",
    "Die folgenden acht Themen sind dabei immer wieder Teil der Prüfung. Geh die Überschriften durch und "
    "entscheide je Thema, wie schwer es in deinem Fall wiegt. Darunter steht, wann das typischerweise "
    "der Fall ist.",
    "Setze den Schwerpunkt bewusst: Prüfungstiefe kostet Geld und Zeit, und beides fehlt dir an anderer "
    "Stelle.",
    "Für Markt-, Technik-, Versicherungs- und Umweltfragen brauchst du je nach Fall einen Spezialisten. "
    "Dein Steuerberater und dein Anwalt aus der Kernprüfung decken diese Themen nicht ab.",
]


# ------------------------------------------------------------------ AutoFit
# Excel misst die noetigen Zeilenhoehen selbst. Verbundene Zellen ignoriert
# Excels AutoFit, deshalb wird ihr Text in eine Hilfsspalte gleicher Breite
# gespiegelt und die Zeile darueber gemessen.
AUTOFIT_PS = r"""
param([string]$Path)
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
$wb = $xl.Workbooks.Open($Path)
$HELP = 26
foreach ($ws in $wb.Worksheets) {
    $used = $ws.UsedRange
    $last = $used.Row + $used.Rows.Count - 1
    $minH = @{}
    for ($r = 1; $r -le $last; $r++) { $minH[$r] = $ws.Rows.Item($r).RowHeight }
    for ($r = 1; $r -le $last; $r++) {
        $text = $null; $fromCol = 0; $span = 0
        for ($c = 2; $c -le 6; $c++) {
            $cell = $ws.Cells.Item($r, $c)
            if ($cell.MergeCells -and $cell.WrapText) {
                $ma = $cell.MergeArea
                if ($ma.Rows.Count -eq 1 -and $ma.Row -eq $r) {
                    $tl = $ws.Cells.Item($r, $ma.Column)
                    if ($tl.Value2 -is [string] -and $tl.Value2.Length -gt 0) {
                        $text = $tl.Value2; $fromCol = $ma.Column; $span = $ma.Columns.Count
                    }
                }
                break
            }
        }
        if ($null -ne $text) {
            $w = 0
            for ($k = 0; $k -lt $span; $k++) { $w += $ws.Columns.Item($fromCol + $k).ColumnWidth }
            $h = $ws.Cells.Item($r, $fromCol)
            $ws.Columns.Item($HELP).ColumnWidth = $w
            $probe = $ws.Cells.Item($r, $HELP)
            $probe.Font.Name = $h.Font.Name; $probe.Font.Size = $h.Font.Size
            $probe.Font.Bold = $h.Font.Bold; $probe.WrapText = $true
            $probe.Value2 = $text
        }
        $fit = ($null -ne $text)
        if (-not $fit) {
            for ($c = 2; $c -le 6; $c++) {
                $cell = $ws.Cells.Item($r, $c)
                if ((-not $cell.MergeCells) -and $cell.WrapText -and $null -ne $cell.Value2) { $fit = $true; break }
            }
        }
        if ($fit) {
            $ws.Rows.Item($r).AutoFit() | Out-Null
            $ws.Rows.Item($r).RowHeight = [Math]::Max($ws.Rows.Item($r).RowHeight + 3, $minH[$r])
            if ($null -ne $text) { $ws.Cells.Item($r, $HELP).Clear() | Out-Null }
        }
    }
    $ws.Columns.Item($HELP).ColumnWidth = 8.43
}
$wb.Worksheets.Item(1).Activate()
$wb.Save(); $wb.Close($true); $xl.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null
"""


def autofit(path):
    """Zeilenhoehen von Excel messen lassen (nur Windows mit installiertem Excel)."""
    import subprocess, tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as fh:
        fh.write(AUTOFIT_PS)
        script = fh.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", script, "-Path", os.path.abspath(path)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            print("AutoFit uebersprungen (kein Excel?):", (r.stderr or "").strip()[:200])
        else:
            print("Zeilenhoehen von Excel gemessen.")
    finally:
        os.unlink(script)




def main():
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Kernprüfung"
    Sheet(ws1, "DD-Checkliste: Kernprüfung",
          "Finanzen, Steuern und Recht. Diese drei prüfst du bei jedem Deal.",
          KERN, ANLEITUNG_KERN).build()

    ws2 = wb.create_sheet("Zusatzbausteine")
    Sheet(ws2, "DD-Checkliste: Zusatzbausteine",
          "Jedes Unternehmen bringt eigene Chancen und Risiken mit. Sie bestimmen Schwerpunkt und Umfang deiner Prüfung.",
          ZUSATZ, ANLEITUNG_ZUSATZ).build()
    ws2.sheet_properties.tabColor = ORANGE

    wb.properties.title = "DUB DD-Checkliste"
    wb.properties.creator = "Deutsche Unternehmerbörse DUB.de GmbH"
    wb.properties.description = "Käufer-Academy, Modul 6: Due Diligence meistern"

    out = os.path.join("Downloads", "Kaeufer", "dd-checkliste.xlsx")
    wb.save(out)
    print("geschrieben:", out)
    autofit(out)


if __name__ == "__main__":
    main()
