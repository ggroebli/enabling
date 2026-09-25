# -*- coding: utf-8 -*-
"""
DUB Kaeufer-Academy - Finanzierungsplan und Kapitaldienst (Modul 7)
Erzeugt Downloads/Kaeufer/finanzierungsplan-kapitaldienst.xlsx

Zwei Blaetter:
  "Finanzierungsplan" - Kapitalbedarf und Mittelherkunft, dazu die
                        Eigenkapitalquote nach der Anrechnungsregel der Bank.
  "Kapitaldienst"     - verfuegbarer Cashflow gegen Annuitaet, mit Stressfall.

Inhaltliche Grundlage: die Leitfaeden des Akquisitionsfinanzierungs-Teams
und Modul 7 der Kaeufer-Academy.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.pagebreak import Break

NAVY      = "0F2744"
ORANGE    = "EE7A1E"
ORANGE_LT = "FDEBD8"
INK       = "1A1A1A"
GREY      = "666666"
GREY_LT   = "8A94A6"
LINE      = "E5E7EB"
PAPER     = "EEF1F5"
INPUT_BG  = "FFF8EF"
INPUT_FG  = "B35C0F"
WHITE     = "FFFFFF"
NAVY_TEXT = "C7D0DE"
FONT = "Calibri"

EUR  = '#,##0 "€";[Red]-#,##0 "€"'
EURC = '#,##0 "€";[Red]-#,##0 "€";""'
PCT  = '0 %'
PCT1 = '0.0 %'
FAK  = '0.00'
JAHR = '0'


def f(sz=10.5, b=False, color=INK, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
LEFTW  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LABELW = Alignment(horizontal="left",   vertical="center", wrap_text=True)


class Blatt:
    """Schreibt ein Rechenblatt zeilenweise und merkt sich die Zeilennummern."""

    def __init__(self, ws, titel, untertitel):
        self.ws = ws
        self.titel = titel
        self.untertitel = untertitel
        self.r = 1
        self.ref = {}
        self.breaks = []

    # -- Bausteine ------------------------------------------------
    def row_h(self, h):
        self.ws.row_dimensions[self.r].height = h

    def band(self, h):
        for col in "ABCDE":
            self.ws[f"{col}{self.r}"].fill = fill(NAVY)
        self.row_h(h)
        self.r += 1

    def section(self, nr, titel):
        ws = self.ws
        self.breaks.append(self.r)
        self.row_h(7)
        self.r += 1
        ws.merge_cells(f"B{self.r}:D{self.r}")
        c = ws[f"B{self.r}"]
        c.value = f"  {nr}   {titel.upper()}"
        c.font = f(11, True, WHITE)
        c.alignment = LEFT
        for col in "BCD":
            ws[f"{col}{self.r}"].fill = fill(NAVY)
        ws[f"A{self.r}"].fill = fill(ORANGE)
        self.row_h(24)
        self.r += 1
        self.row_h(5)          # weisse Trennzeile, damit Eingabefelder nicht ans Band stossen
        self.r += 1

    def note(self, text, color=GREY, size=9.5, italic=False, bg=None):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:D{self.r}")
        c = ws[f"B{self.r}"]
        c.value = text
        c.font = f(size, False, color, italic)
        c.alignment = LEFTW
        if bg:
            for col in "BCD":
                ws[f"{col}{self.r}"].fill = fill(bg)
        self.row_h(14)
        self.r += 1

    def eingabe(self, label, wert=None, fmt=EUR, hinweis="", indent=False):
        ws = self.ws
        c = ws[f"B{self.r}"]
        c.value = ("      " if indent else "") + label
        c.font = f(10.5)
        c.alignment = LABELW
        c.border = Border(bottom=side(LINE))
        z = ws[f"C{self.r}"]
        z.value = wert
        z.number_format = fmt
        z.font = f(10.5, False, INPUT_FG)
        z.fill = fill(INPUT_BG)
        z.alignment = RIGHT
        z.border = Border(bottom=side(LINE), left=side(WHITE, "thin"))
        h = ws[f"D{self.r}"]
        h.value = hinweis
        h.font = f(9.5, False, GREY_LT)
        h.alignment = LEFTW
        h.border = Border(bottom=side(LINE))
        self.row_h(18)
        row = self.r
        self.r += 1
        return row

    def rechnung(self, label, formel, fmt=EURC, akzent=False, size=10.5, hinweis="", height=19):
        ws = self.ws
        bg = ORANGE_LT if akzent else PAPER
        c = ws[f"B{self.r}"]
        c.value = label
        if isinstance(label, str) and label.startswith("="):
            c.data_type = "s"
        c.font = f(size, True, NAVY if akzent else INK)
        c.alignment = LABELW
        c.fill = fill(bg)
        c.border = Border(top=side(NAVY if akzent else GREY_LT), bottom=side(LINE))
        z = ws[f"C{self.r}"]
        z.value = formel
        z.number_format = fmt
        z.font = f(size, True, NAVY if akzent else INK)
        z.fill = fill(bg)
        z.alignment = RIGHT
        z.border = Border(top=side(NAVY if akzent else GREY_LT), bottom=side(LINE))
        h = ws[f"D{self.r}"]
        h.value = hinweis
        h.font = f(9.5, False, GREY_LT)
        h.alignment = LEFTW
        h.fill = fill(bg)
        h.border = Border(top=side(NAVY if akzent else GREY_LT), bottom=side(LINE))
        self.row_h(height)
        row = self.r
        self.r += 1
        return row

    def spacer(self, h=8):
        self.row_h(h)
        self.r += 1

    def kopf(self):
        ws = self.ws
        self.band(6)
        ws[f"B{self.r}"] = "DUB  ·  KÄUFER-ACADEMY"
        ws[f"B{self.r}"].font = f(10, True, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        ws[f"D{self.r}"] = "MODUL 7 · AKQUISITIONSFINANZIERUNG"
        ws[f"D{self.r}"].font = f(10, True, ORANGE)
        ws[f"D{self.r}"].alignment = Alignment(horizontal="right", vertical="center")
        self.band(18)
        ws.merge_cells(f"B{self.r}:D{self.r}")
        ws[f"B{self.r}"] = self.titel
        ws[f"B{self.r}"].font = f(20, True, WHITE)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(30)
        ws.merge_cells(f"B{self.r}:D{self.r}")
        ws[f"B{self.r}"] = self.untertitel
        ws[f"B{self.r}"].font = f(11, False, NAVY_TEXT)
        ws[f"B{self.r}"].alignment = LEFT
        self.band(19)
        self.band(8)
        self.spacer(10)

    def fuss(self):
        ws = self.ws
        ws.merge_cells(f"B{self.r}:D{self.r}")
        c = ws[f"B{self.r}"]
        c.value = ("Diese Vorlage dient der Orientierung und ersetzt keine Rechts-, Steuer- oder "
                   "Finanzberatung.    DUB · info@dub.de · www.dub.de    "
                   "© 2026 Deutsche Unternehmerbörse DUB.de GmbH")
        c.font = f(9, False, GREY_LT)
        c.alignment = LEFTW
        for col in "BCD":
            ws[f"{col}{self.r}"].border = Border(top=side(LINE))
        self.row_h(24)
        self.r += 1

    def layout(self):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = NAVY
        for col, w in {"A": 2.2, "B": 62, "C": 18, "D": 46, "E": 2.2}.items():
            ws.column_dimensions[col].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4


def blatt_plan(ws):
    b = Blatt(ws, "Finanzierungsplan",
              "Was am Tag 1 bezahlt werden muss – und woher das Geld kommt")
    b.layout()
    b.kopf()

    b.section("▶", "So gehst du vor")
    for t in [
        "1.  Trage in Abschnitt 1 deinen Kapitalbedarf ein. Der Kaufpreis ist nur ein Teil davon: "
        "Beratung, Notar und eine Reserve für die ersten Monate kommen dazu.",
        "2.  Verteile in Abschnitt 2 den Bedarf auf die Bausteine. Die Prüfzeile zeigt dir, ob die "
        "Rechnung aufgeht.",
        "3.  Abschnitt 3 rechnet aus, was die Bank davon als Eigenkapital anerkennt. Das ist fast immer "
        "weniger, als du selbst zusammenzählst.",
    ]:
        b.note(t, INK, 10)
    b.note("Die orange hinterlegten Felder füllst du aus. Alles andere rechnet sich.", GREY_LT, 9.5, italic=True)
    b.spacer(10)

    # -- 1 Kapitalbedarf
    b.section("1", "Kapitalbedarf")
    start = b.r
    b.eingabe("Kaufpreis", hinweis="Ergebnis deiner Preisbrücke aus dem Bewertungs-Template")
    b.eingabe("Beratung: Anwalt, Steuerberater, Prüfung", hinweis="Bei kleinen Deals grob 5 bis 10 % des Kaufpreises")
    b.eingabe("Notar und Registerkosten", hinweis="Beim Share Deal Pflicht, beim Asset Deal je nach Gegenstand")
    b.eingabe("Grunderwerbsteuer", hinweis="Nur wenn eine Immobilie mit übergeht")
    b.eingabe("Liquiditätsreserve für die ersten Monate", hinweis="Puffer für Anlaufkosten und schwankende Zahlungseingänge")
    b.eingabe("Weiterer Bedarf (frei)")
    ende = b.r - 1
    b.ref["bedarf"] = b.rechnung("= GESAMTER KAPITALBEDARF", f"=SUM(C{start}:C{ende})",
                                 akzent=True, size=12, height=24,
                                 hinweis="Auf diese Summe bezieht die Bank deine Eigenkapitalquote")
    b.spacer(10)

    # -- 2 Mittelherkunft
    b.section("2", "Woher das Geld kommt")
    s2 = b.r
    b.ref["ek"] = b.eingabe("Echtes Eigenkapital", hinweis="Konto, Wertpapiere, Mitinvestoren, Familiendarlehen mit Rangrücktritt")
    b.ref["vl"] = b.eingabe("Verkäuferdarlehen, nachrangig", hinweis="Nur mit qualifiziertem Rangrücktritt und tilgungsfreien Anfangsjahren")
    b.ref["mezz"] = b.eingabe("Mezzanine oder stille Beteiligung", hinweis="Zum Beispiel über die Beteiligungsgesellschaft deines Bundeslandes")
    b.ref["bank"] = b.eingabe("Bankdarlehen", hinweis="Der Rest. Diese Zahl trägt Blatt 2 in die Kapitaldienstrechnung")
    e2 = b.r - 1
    b.ref["summe"] = b.rechnung("= Summe der Mittel", f"=SUM(C{s2}:C{e2})")
    zeile = b.rechnung("Differenz zum Kapitalbedarf",
                       "=C{s}-C{b}".format(s=b.ref["summe"], b=b.ref["bedarf"]),
                       fmt=EUR)
    ws[f"D{zeile}"] = ('=IF(C{b}=0,"",IF(C{z}=0,"Passt, der Plan geht auf.",'
                       'IF(C{z}<0,"Es fehlen noch "&TEXT(-C{z},"#.##0")&" €.",'
                       '"Du hast "&TEXT(C{z},"#.##0")&" € zu viel eingeplant.")))'
                       ).format(z=zeile, b=b.ref["bedarf"])
    ws[f"D{zeile}"].font = f(9.5, True, NAVY)
    b.spacer(10)

    # -- 3 Anerkanntes Eigenkapital
    b.section("3", "Was die Bank als Eigenkapital anerkennt")
    b.note("Ein nachrangiges Verkäuferdarlehen zählt nur anteilig, in der Praxis häufig zur Hälfte – und "
           "höchstens bis zur Höhe deines echten Eigenkapitals. Mezzanine bleibt Fremdkapital, auch wenn es "
           "nachrangig ist.", INK, 10)
    b.ref["anr"] = b.rechnung("Anrechenbarer Teil des Verkäuferdarlehens",
                              "=MIN(C{v}/2,C{e})".format(v=b.ref["vl"], e=b.ref["ek"]),
                              hinweis="Die Hälfte, gedeckelt auf dein echtes Eigenkapital")
    b.ref["wirt"] = b.rechnung("= Wirtschaftliches Eigenkapital",
                               "=C{e}+C{a}".format(e=b.ref["ek"], a=b.ref["anr"]),
                               akzent=True, size=12, height=24)
    b.ref["quote"] = b.rechnung("Quote am Kapitalbedarf",
                                '=IFERROR(C{w}/C{b},"")'.format(w=b.ref["wirt"], b=b.ref["bedarf"]),
                                fmt=PCT1, size=11,
                                hinweis="Banken erwarten mindestens 10 bis 20 %, ohne Branchenerfahrung eher mehr")
    ws.merge_cells(f"B{b.r}:D{b.r}")
    c = ws[f"B{b.r}"]
    c.value = ('=IF(C{q}="","",IF(C{q}<0.1,"Unter 10 %: Für die meisten Banken zu wenig. Mehr Eigenkapital '
               'oder ein niedrigerer Kaufpreis.",IF(C{q}<0.2,"10 bis 20 %: Machbar, aber ohne Puffer. Die '
               'Struktur muss überzeugen.","Über 20 %: Komfortabler Bereich, du kannst über Konditionen '
               'verhandeln.")))').format(q=b.ref["quote"])
    c.font = f(10, True, NAVY)
    c.alignment = LEFTW
    for col in "BCD":
        ws[f"{col}{b.r}"].fill = fill(ORANGE_LT)
    b.row_h(17)
    b.r += 1
    b.spacer(10)
    b.fuss()
    for r in b.breaks[2:]:
        ws.row_breaks.append(Break(id=r))
    return b


def blatt_dienst(ws, plan):
    b = Blatt(ws, "Kapitaldienst",
              "Verdient das Unternehmen genug, um dein Darlehen zu bedienen?")
    b.layout()
    ws.sheet_properties.tabColor = ORANGE
    b.kopf()

    b.section("▶", "So gehst du vor")
    for t in [
        "Das Bankdarlehen kommt automatisch aus Blatt 1. Ergänze hier die Ertragskraft des Unternehmens "
        "und die Konditionen, die deine Bank nennt.",
        "Die Kennzahl unten ist die eine Rechnung, die jede Bank macht. Wer sie mitbringt, führt das "
        "Gespräch, statt es zu erleben.",
    ]:
        b.note(t, INK, 10)
    b.spacer(10)

    # -- 1 Ertragskraft
    b.section("1", "Was das Unternehmen erwirtschaftet")
    b.ref["ebitda"] = b.eingabe("Bereinigtes EBITDA", hinweis="Aus deinem Bewertungs-Template, Abschnitt 2")
    b.ref["anteil"] = b.eingabe("Davon verfügbar für den Kapitaldienst", 0.70, fmt=PCT,
                                hinweis="Faustregel 70 %. Produktion und Logistik eher weniger, Dienstleister mehr")
    b.note("Der Rest wird gebraucht für Steuern, Erhaltungsinvestitionen, gebundenes Kapital im "
           "Tagesgeschäft und ein angemessenes Geschäftsführergehalt.", GREY_LT, 9.5, italic=True)
    b.ref["cf"] = b.rechnung("= Verfügbarer Cashflow",
                             "=C{e}*C{a}".format(e=b.ref["ebitda"], a=b.ref["anteil"]),
                             akzent=True, size=12, height=24)
    b.spacer(10)

    # -- 2 Kapitaldienst
    b.section("2", "Was der Kredit jedes Jahr kostet")
    b.ref["bank"] = b.eingabe("Bankdarlehen", f"='Finanzierungsplan'!C{plan.ref['bank']}",
                              hinweis="Kommt aus Blatt 1, Abschnitt 2")
    b.ref["zins"] = b.eingabe("Zinssatz", 0.06, fmt=PCT1, hinweis="Frag deine Bank. Ohne Angabe ist das hier nur eine Übung")
    b.ref["jahre"] = b.eingabe("Laufzeit in Jahren", 7, fmt=JAHR,
                               hinweis="Marktstandard sechs bis sieben. Wer zehn braucht, zahlt zu viel")
    b.ref["ann"] = b.rechnung("Jährliche Annuität für das Bankdarlehen",
                              '=IFERROR(-PMT(C{z},C{j},C{b}),"")'.format(z=b.ref["zins"], j=b.ref["jahre"], b=b.ref["bank"]),
                              hinweis="Zins und Tilgung in gleichbleibenden Raten")
    b.ref["nach"] = b.eingabe("Zinsen auf Verkäuferdarlehen und Mezzanine",
                              hinweis="Nur die Zinsen. Die Tilgung folgt erst nach der Bankphase")
    b.ref["kd"] = b.rechnung("= Gesamter Kapitaldienst pro Jahr",
                             "=C{a}+C{n}".format(a=b.ref["ann"], n=b.ref["nach"]),
                             akzent=True, size=12, height=24)
    b.spacer(10)

    # -- 3 Ergebnis
    b.section("3", "Das Ergebnis")
    b.ref["dscr"] = b.rechnung("Kapitaldienstfähigkeit",
                               '=IFERROR(C{c}/C{k},"")'.format(c=b.ref["cf"], k=b.ref["kd"]),
                               fmt=FAK, akzent=True, size=14, height=28,
                               hinweis="Verfügbarer Cashflow geteilt durch Kapitaldienst")
    ws.merge_cells(f"B{b.r}:D{b.r}")
    c = ws[f"B{b.r}"]
    c.value = ('=IF(C{d}="","",IF(C{d}<1.2,"Unter 1,2: Nicht tragfähig. Kaufpreis senken, Struktur ändern '
               'oder Laufzeit verlängern.",IF(C{d}<1.5,"1,2 bis 1,5: Machbar, aber die Bank wird Anpassungen '
               'verlangen.","Ab 1,5: Komfortabel finanzierbar. Du kannst über Konditionen verhandeln.")))'
               ).format(d=b.ref["dscr"])
    c.font = f(10, True, NAVY)
    c.alignment = LEFTW
    for col in "BCD":
        ws[f"{col}{b.r}"].fill = fill(ORANGE_LT)
    b.row_h(17)
    b.r += 1
    b.rechnung("Verschuldung im Verhältnis zum EBITDA",
               '=IFERROR((C{b}+{nl})/C{e},"")'.format(b=b.ref["bank"],
                   nl="'Finanzierungsplan'!C%d+'Finanzierungsplan'!C%d" % (plan.ref["vl"], plan.ref["mezz"]),
                   e=b.ref["ebitda"]),
               fmt=FAK, size=10,
               hinweis="Alle Darlehen zusammen. Wie viele Jahresgewinne bis zur Schuldenfreiheit?")
    b.spacer(10)

    # -- 4 Stressfall
    b.section("4", "Wenn es schlechter läuft")
    b.note("Banken rechnen nicht nur den Planfall. Sie senken das Ergebnis und schauen, ab wann die "
           "Rechnung kippt. Genau dieser Puffer ist der Grund, warum sie 1,5 sehen wollen und nicht 1,0.",
           INK, 10)
    for col, t in zip("BCD", ["Szenario", "Kapitaldienstfähigkeit", "Verfügbarer Cashflow"]):
        z = ws[f"{col}{b.r}"]
        z.value = t
        z.font = f(9.5, True, NAVY)
        z.alignment = Alignment(horizontal="right" if col != "B" else "left", vertical="bottom", wrap_text=True)
        z.border = Border(bottom=side(ORANGE, "medium"))
    b.row_h(20)
    b.r += 1
    for minus, label in [(0, "Planfall"), (0.2, "EBITDA minus 20 Prozent"),
                         (0.3, "EBITDA minus 30 Prozent"), (0.4, "EBITDA minus 40 Prozent")]:
        z = ws[f"B{b.r}"]
        z.value = label
        z.font = f(10.5, minus == 0, NAVY if minus == 0 else INK)
        z.alignment = LABELW
        cf = "C{e}*(1-{m})*C{a}".format(e=b.ref["ebitda"], m=minus, a=b.ref["anteil"])
        d = ws[f"C{b.r}"]
        d.value = '=IFERROR({cf}/C{k},"")'.format(cf=cf, k=b.ref["kd"])
        d.number_format = FAK
        d.font = f(10.5, True, INK)
        d.alignment = RIGHT
        e = ws[f"D{b.r}"]
        e.value = "={cf}".format(cf=cf)
        e.number_format = EURC
        e.font = f(10, False, GREY)
        e.alignment = RIGHT
        for col in "BCD":
            ws[f"{col}{b.r}"].border = Border(bottom=side(LINE))
        b.row_h(18)
        b.r += 1
    b.spacer(6)
    b.note("Kippt die Rechnung schon bei minus 20 Prozent, ist der Kaufpreis zu hoch oder die Struktur "
           "falsch. Beides lässt sich vor dem Bankgespräch ändern, danach nicht mehr.", GREY, 9.5)
    b.spacer(10)
    b.fuss()
    for r in b.breaks[2:]:
        ws.row_breaks.append(Break(id=r))
    return b


# ------------------------------------------------------------------ AutoFit
# Excel misst die noetigen Zeilenhoehen selbst. Verbundene Zellen ignoriert
# Excels AutoFit, deshalb wird ihr Text in eine Hilfsspalte gleicher Breite
# gespiegelt und die Zeile darueber gemessen.
AUTOFIT_PS = r"""
param([string]$Path)
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
$wb = $xl.Workbooks.Open($Path)
$HELP = 26
foreach ($ws in $wb.Worksheets) {
    $used = $ws.UsedRange
    $last = $used.Row + $used.Rows.Count - 1
    $minH = @{}
    for ($r = 1; $r -le $last; $r++) { $minH[$r] = $ws.Rows.Item($r).RowHeight }
    for ($r = 1; $r -le $last; $r++) {
        $text = $null; $fromCol = 0; $span = 0
        for ($c = 2; $c -le 6; $c++) {
            $cell = $ws.Cells.Item($r, $c)
            if ($cell.MergeCells -and $cell.WrapText) {
                $ma = $cell.MergeArea
                if ($ma.Rows.Count -eq 1 -and $ma.Row -eq $r) {
                    $tl = $ws.Cells.Item($r, $ma.Column)
                    if ($tl.Value2 -is [string] -and $tl.Value2.Length -gt 0) {
                        $text = $tl.Value2; $fromCol = $ma.Column; $span = $ma.Columns.Count
                    }
                }
                break
            }
        }
        if ($null -ne $text) {
            $w = 0
            for ($k = 0; $k -lt $span; $k++) { $w += $ws.Columns.Item($fromCol + $k).ColumnWidth }
            $h = $ws.Cells.Item($r, $fromCol)
            $ws.Columns.Item($HELP).ColumnWidth = $w
            $probe = $ws.Cells.Item($r, $HELP)
            $probe.Font.Name = $h.Font.Name; $probe.Font.Size = $h.Font.Size
            $probe.Font.Bold = $h.Font.Bold; $probe.WrapText = $true
            $probe.Value2 = $text
        }
        $fit = ($null -ne $text)
        if (-not $fit) {
            for ($c = 2; $c -le 6; $c++) {
                $cell = $ws.Cells.Item($r, $c)
                if ((-not $cell.MergeCells) -and $cell.WrapText -and $null -ne $cell.Value2) { $fit = $true; break }
            }
        }
        if ($fit) {
            $ws.Rows.Item($r).AutoFit() | Out-Null
            $ws.Rows.Item($r).RowHeight = [Math]::Max($ws.Rows.Item($r).RowHeight + 3, $minH[$r])
            if ($null -ne $text) { $ws.Cells.Item($r, $HELP).Clear() | Out-Null }
        }
    }
    $ws.Columns.Item($HELP).ColumnWidth = 8.43
}
$wb.Worksheets.Item(1).Activate()
$wb.Save(); $wb.Close($true); $xl.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null
"""


def autofit(path):
    """Zeilenhoehen von Excel messen lassen (nur Windows mit installiertem Excel)."""
    import subprocess, tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as fh:
        fh.write(AUTOFIT_PS)
        script = fh.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", script, "-Path", os.path.abspath(path)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            print("AutoFit uebersprungen (kein Excel?):", (r.stderr or "").strip()[:200])
        else:
            print("Zeilenhoehen von Excel gemessen.")
    finally:
        os.unlink(script)




def main():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Finanzierungsplan"
    plan = blatt_plan(ws1)

    ws2 = wb.create_sheet("Kapitaldienst")
    blatt_dienst(ws2, plan)

    wb.properties.title = "DUB Finanzierungsplan und Kapitaldienst"
    wb.properties.creator = "Deutsche Unternehmerboerse DUB.de GmbH"
    wb.properties.description = "Kaeufer-Academy, Modul 7: Akquisitionsfinanzierung"

    out = os.path.join("Downloads", "Kaeufer", "finanzierungsplan-kapitaldienst.xlsx")
    wb.save(out)
    print("geschrieben:", out)
    autofit(out)


if __name__ == "__main__":
    main()
